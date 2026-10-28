import os
import io
import re
import json
import uuid
import sqlite3
import base64
import threading
import urllib.request
from datetime import datetime, date, time
from functools import wraps
import numpy as np
import cv2
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from flask import (
    Flask, render_template, request, redirect, url_for,
    session, jsonify, send_file
)
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from fpdf import FPDF

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', '9f8d7a6b5c4e3f2a1b0c9d8e7f6a5b4c3d2e1f0a9b8c7d6e5f4a3b2c1d0e9f8a')

app.config['SESSION_COOKIE_HTTPONLY'] = True
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'
app.config['MAX_CONTENT_LENGTH'] = 25 * 1024 * 1024

BASE_DIR = os.path.abspath(os.path.dirname(__file__))
DB_PATH = os.path.join(BASE_DIR, 'refeitorio.db')
UPLOAD_FOLDER = os.path.join(BASE_DIR, 'static', 'uploads')
MODELS_FOLDER = os.path.join(BASE_DIR, 'models')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(MODELS_FOLDER, exist_ok=True)

ALLOWED_EXTENSIONS = {'png', 'jpg', 'jpeg', 'webp'}

YUNET_PATH = os.path.join(MODELS_FOLDER, 'face_detection_yunet_2023mar.onnx')
SFACE_PATH = os.path.join(MODELS_FOLDER, 'face_recognition_sface_2021dec.onnx')

YUNET_URL = 'https://huggingface.co/opencv/face_detection_yunet/resolve/main/face_detection_yunet_2023mar.onnx'
SFACE_URL = 'https://huggingface.co/opencv/face_recognition_sface/resolve/main/face_recognition_sface_2021dec.onnx'

LIMIAR_SIMILARIDADE = 0.38
MAX_VETORES_POR_COLABORADOR = 15

FACIAL_LOCK = threading.Lock()
_detector_yn = None
_recognizer_sf = None

TEMAS = {
    'tema_1': {'nome': 'Azul', 'primaria': '#1e3a8a', 'secundaria': '#2563eb', 'fundo': '#f8fafc', 'texto': '#0f172a'},
    'tema_2': {'nome': 'Verde', 'primaria': '#065f46', 'secundaria': '#10b981', 'fundo': '#f0fdf4', 'texto': '#064e3b'},
    'tema_3': {'nome': 'Roxo', 'primaria': '#4c1d95', 'secundaria': '#7c3aed', 'fundo': '#faf5ff', 'texto': '#2e1065'},
    'tema_4': {'nome': 'Grafite', 'primaria': '#18181b', 'secundaria': '#52525b', 'fundo': '#f4f4f5', 'texto': '#09090b'},
    'tema_5': {'nome': 'Índigo', 'primaria': '#312e81', 'secundaria': '#4f46e5', 'fundo': '#eef2ff', 'texto': '#1e1b4b'},
    'tema_6': {'nome': 'Turquesa', 'primaria': '#115e59', 'secundaria': '#0d9488', 'fundo': '#f0fdfa', 'texto': '#134e4a'},
    'tema_7': {'nome': 'Rubi', 'primaria': '#881337', 'secundaria': '#e11d48', 'fundo': '#fff1f2', 'texto': '#4c0519'},
    'tema_8': {'nome': 'Âmbar', 'primaria': '#78350f', 'secundaria': '#d97706', 'fundo': '#fffbeb', 'texto': '#451a03'},
    'tema_9': {'nome': 'Azul Petróleo', 'primaria': '#164e63', 'secundaria': '#0891b2', 'fundo': '#ecfeff', 'texto': '#083344'},
    'tema_10': {'nome': 'Lavanda', 'primaria': '#581c87', 'secundaria': '#9333ea', 'fundo': '#fdf4ff', 'texto': '#3b0764'}
}

# Mapeamento de abas e permissões correspondentes
MAPA_ABAS_PERMISSOES = [
    ('dashboard', 'dashboard'),
    ('usuarios', 'cad_usuario'),
    ('refeicoes', 'cad_refeicao'),
    ('grupos', 'cad_grupo'),
    ('colaboradores', 'cad_colaborador'),
    ('cardapios', 'cad_cardapio'),
    ('unidades', 'cad_unidade'),
    ('importacao', 'importacao'),
    ('relatorios', 'relatorios'),
    ('logs', 'logs'),
    ('configuracoes', 'configuracoes')
]

def obter_aba_inicial_permitida(permissoes):
    for nome_aba, perm_necessaria in MAPA_ABAS_PERMISSOES:
        if perm_necessaria in permissoes:
            return nome_aba
    return 'sem_acesso'

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn

def limpar_cpf_cnpj(valor):
    return re.sub(r'\D', '', str(valor or ''))

def arquivo_permitido(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def usuario_eh_multiempresa():
    return 'multiempresa' in session.get('permissoes', [])

# ==============================================================================
# MOTOR DE INTELIGÊNCIA ARTIFICIAL FACIAL (YUNET + SFACE + CLAHE)
# ==============================================================================
def baixar_modelo_se_necessario(url, destino, tamanho_minimo_bytes):
    if os.path.exists(destino) and os.path.getsize(destino) >= tamanho_minimo_bytes:
        return True
    tmp_path = f"{destino}.{uuid.uuid4().hex}.tmp"
    try:
        print(f"Baixando modelo facial: {os.path.basename(destino)}...")
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=120) as response, open(tmp_path, 'wb') as out_file:
            out_file.write(response.read())
        if os.path.getsize(tmp_path) >= tamanho_minimo_bytes:
            os.replace(tmp_path, destino)
            print(f"Modelo {os.path.basename(destino)} pronto.")
            return True
    except Exception as e:
        print(f"Erro ao baixar modelo {destino}: {e}")
    finally:
        if os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except Exception:
                pass
    return False

def obter_motores_faciais():
    global _detector_yn, _recognizer_sf
    if _detector_yn is not None and _recognizer_sf is not None:
        return _detector_yn, _recognizer_sf

    ok1 = baixar_modelo_se_necessario(YUNET_URL, YUNET_PATH, 100_000)
    ok2 = baixar_modelo_se_necessario(SFACE_URL, SFACE_PATH, 10_000_000)
    if not (ok1 and ok2):
        return None, None

    _detector_yn = cv2.FaceDetectorYN.create(
        YUNET_PATH, "", (320, 320),
        score_threshold=0.6, nms_threshold=0.3, top_k=5000
    )
    _recognizer_sf = cv2.FaceRecognizerSF.create(SFACE_PATH, "")
    return _detector_yn, _recognizer_sf

def corrigir_iluminacao_clahe(img_bgr):
    try:
        lab = cv2.cvtColor(img_bgr, cv2.COLOR_BGR2LAB)
        l, a, b = cv2.split(lab)
        clahe = cv2.createCLAHE(clipLimit=2.5, tileGridSize=(8, 8))
        l_eq = clahe.apply(l)
        lab_eq = cv2.merge((l_eq, a, b))
        return cv2.cvtColor(lab_eq, cv2.COLOR_LAB2BGR)
    except Exception:
        return img_bgr

def decodificar_base64_para_bgr(b64_str):
    if not b64_str or ',' not in b64_str:
        return None
    try:
        _, encoded = b64_str.split(',', 1)
        dados = base64.b64decode(encoded)
        arr = np.frombuffer(dados, dtype=np.uint8)
        img = cv2.imdecode(arr, cv2.IMREAD_COLOR)
        return img
    except Exception:
        return None

def analisar_rosto_bgr(img_bgr):
    if img_bgr is None or img_bgr.size == 0:
        return None, None, 0.0

    with FACIAL_LOCK:
        detector, recognizer = obter_motores_faciais()
        if detector is None or recognizer is None:
            return None, None, 0.0

        img_corrigida = corrigir_iluminacao_clahe(img_bgr)
        h, w = img_corrigida.shape[:2]

        escala = 1.0
        largura_max = 640
        if w > largura_max:
            escala = largura_max / float(w)
            nova_w = int(w * escala)
            nova_h = int(h * escala)
            img_proc = cv2.resize(img_corrigida, (nova_w, nova_h))
        else:
            img_proc = img_corrigida
            nova_w, nova_h = w, h

        detector.setInputSize((nova_w, nova_h))
        _, faces = detector.detect(img_proc)

        if faces is None or len(faces) == 0:
            return None, None, 0.0

        maior_face = max(faces, key=lambda f: f[2] * f[3])
        fx, fy, fw, fh = [float(v) for v in maior_face[:4]]

        proporcao = (fw * fh) / float(nova_w * nova_h)
        bbox_rel = {
            'x': max(0.0, fx / nova_w),
            'y': max(0.0, fy / nova_h),
            'w': min(1.0, fw / nova_w),
            'h': min(1.0, fh / nova_h)
        }

        rosto_alinhado = recognizer.alignCrop(img_proc, maior_face)
        feat = recognizer.feature(rosto_alinhado)
        vetor = feat.flatten().astype(np.float32)
        norma = np.linalg.norm(vetor)
        if norma > 0:
            vetor = vetor / norma

        return vetor.tolist(), bbox_rel, proporcao

def comparar_vetor_com_lista(vetor_consulta, lista_vetores):
    if not vetor_consulta or not lista_vetores:
        return 0.0
    v_query = np.array(vetor_consulta, dtype=np.float32)
    melhor_score = 0.0
    for v_salvo in lista_vetores:
        v_db = np.array(v_salvo, dtype=np.float32)
        if v_db.shape == v_query.shape:
            score = float(np.dot(v_query, v_db))
            if score > melhor_score:
                melhor_score = score
    return melhor_score

def adicionar_vetor_aprendizado_continuo(colaborador_id, novo_vetor):
    if not novo_vetor:
        return
    try:
        conn = get_db()
        row = conn.execute('SELECT face_embeddings FROM colaboradores WHERE id = ?', (colaborador_id,)).fetchone()
        if not row:
            conn.close()
            return
        atuais = json.loads(row['face_embeddings'] or '[]')

        melhor_atual = comparar_vetor_com_lista(novo_vetor, atuais)
        if melhor_atual > 0.82 and len(atuais) >= 3:
            conn.close()
            return

        if len(atuais) < MAX_VETORES_POR_COLABORADOR:
            atuais.append(novo_vetor)
        else:
            originais = atuais[:5]
            aprendidos = atuais[5:]
            aprendidos.append(novo_vetor)
            atuais = originais + aprendidos[-(MAX_VETORES_POR_COLABORADOR - 5):]

        conn.execute('UPDATE colaboradores SET face_embeddings = ? WHERE id = ?', (json.dumps(atuais), colaborador_id))
        conn.commit()
        conn.close()
    except Exception as e:
        print(f"Erro no aprendizado contínuo: {e}")

# ==============================================================================
# SEGURANÇA E INICIALIZAÇÃO DO BANCO DE DADOS
# ==============================================================================
@app.after_request
def aplicar_cabecalhos_seguranca(response):
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['X-Frame-Options'] = 'SAMEORIGIN'
    response.headers['X-XSS-Protection'] = '1; mode=block'
    response.headers['Referrer-Policy'] = 'strict-origin-when-cross-origin'
    return response

@app.before_request
def verificar_csrf():
    if 'csrf_token' not in session:
        session['csrf_token'] = uuid.uuid4().hex
    if request.method == 'POST':
        token = request.headers.get('X-CSRFToken') or request.form.get('csrf_token')
        if not token or token != session.get('csrf_token'):
            return jsonify({'status': 'erro', 'mensagem': 'Sessão expirada. Recarregue a página.'}), 403

def registrar_log(tabela, acao, descricao, unidade_id=None):
    try:
        conn = get_db()
        usuario_id = session.get('usuario_id')
        usuario_nome = session.get('usuario_nome', 'Sistema / Totem')
        uid = unidade_id or session.get('unidade_id') or session.get('totem_unidade_id')
        agora = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        conn.execute('''
            INSERT INTO logs (data_hora, usuario_id, usuario_nome, unidade_id, tabela, acao, descricao)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        ''', (agora, usuario_id, usuario_nome, uid, tabela, acao, descricao))
        conn.commit()
        conn.close()
    except Exception:
        pass

def init_db():
    conn = get_db()
    c = conn.cursor()

    c.execute('''
        CREATE TABLE IF NOT EXISTS unidades (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            descricao TEXT NOT NULL,
            ativo INTEGER DEFAULT 1
        )
    ''')

    c.execute('''
        CREATE TABLE IF NOT EXISTS configuracoes (
            unidade_id INTEGER PRIMARY KEY,
            tema TEXT DEFAULT 'tema_1',
            logo_path TEXT DEFAULT '',
            FOREIGN KEY (unidade_id) REFERENCES unidades(id)
        )
    ''')

    c.execute('''
        CREATE TABLE IF NOT EXISTS usuarios (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nome TEXT NOT NULL,
            cpf TEXT UNIQUE NOT NULL,
            cargo TEXT,
            senha_hash TEXT NOT NULL,
            primeiro_acesso INTEGER DEFAULT 1,
            ativo INTEGER DEFAULT 1,
            permissoes TEXT NOT NULL
        )
    ''')

    c.execute('''
        CREATE TABLE IF NOT EXISTS usuario_unidades (
            usuario_id INTEGER,
            unidade_id INTEGER,
            PRIMARY KEY (usuario_id, unidade_id),
            FOREIGN KEY (usuario_id) REFERENCES usuarios(id) ON DELETE CASCADE,
            FOREIGN KEY (unidade_id) REFERENCES unidades(id) ON DELETE CASCADE
        )
    ''')

    c.execute('''
        CREATE TABLE IF NOT EXISTS grupos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            descricao TEXT NOT NULL,
            ativo INTEGER DEFAULT 1,
            unidade_id INTEGER NOT NULL,
            FOREIGN KEY (unidade_id) REFERENCES unidades(id)
        )
    ''')

    c.execute('''
        CREATE TABLE IF NOT EXISTS refeicoes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            descricao TEXT NOT NULL,
            hora_inicio TEXT NOT NULL,
            hora_fim TEXT NOT NULL,
            custo REAL NOT NULL DEFAULT 0.0,
            ativo INTEGER DEFAULT 1,
            unidade_id INTEGER NOT NULL,
            FOREIGN KEY (unidade_id) REFERENCES unidades(id)
        )
    ''')

    c.execute('''
        CREATE TABLE IF NOT EXISTS refeicao_grupos (
            refeicao_id INTEGER,
            grupo_id INTEGER,
            PRIMARY KEY (refeicao_id, grupo_id),
            FOREIGN KEY (refeicao_id) REFERENCES refeicoes(id) ON DELETE CASCADE,
            FOREIGN KEY (grupo_id) REFERENCES grupos(id) ON DELETE CASCADE
        )
    ''')

    c.execute('''
        CREATE TABLE IF NOT EXISTS colaboradores (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nome TEXT NOT NULL,
            cpf TEXT NOT NULL,
            foto_path TEXT,
            face_embeddings TEXT DEFAULT '[]',
            grupo_id INTEGER NOT NULL,
            ativo INTEGER DEFAULT 1,
            unidade_id INTEGER NOT NULL,
            UNIQUE(cpf, unidade_id),
            FOREIGN KEY (grupo_id) REFERENCES grupos(id),
            FOREIGN KEY (unidade_id) REFERENCES unidades(id)
        )
    ''')

    colunas_colab = [col['name'] for col in c.execute("PRAGMA table_info(colaboradores)").fetchall()]
    if 'face_embeddings' not in colunas_colab:
        c.execute("ALTER TABLE colaboradores ADD COLUMN face_embeddings TEXT DEFAULT '[]'")

    c.execute('''
        CREATE TABLE IF NOT EXISTS cardapios (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            data_vigencia TEXT NOT NULL,
            hora_inicio TEXT NOT NULL DEFAULT '',
            hora_fim TEXT NOT NULL DEFAULT '',
            guarnicao TEXT NOT NULL,
            proteinas TEXT NOT NULL,
            saladas TEXT NOT NULL,
            sobremesas TEXT NOT NULL,
            bebidas TEXT NOT NULL,
            ativo INTEGER DEFAULT 1,
            unidade_id INTEGER NOT NULL,
            FOREIGN KEY (unidade_id) REFERENCES unidades(id)
        )
    ''')

    c.execute('''
        CREATE TABLE IF NOT EXISTS cardapio_refeicoes (
            cardapio_id INTEGER,
            refeicao_id INTEGER,
            PRIMARY KEY (cardapio_id, refeicao_id),
            FOREIGN KEY (cardapio_id) REFERENCES cardapios(id) ON DELETE CASCADE,
            FOREIGN KEY (refeicao_id) REFERENCES refeicoes(id) ON DELETE CASCADE
        )
    ''')

    cardapios_antigos = c.execute('''
        SELECT c.id, c.unidade_id, c.hora_inicio, c.hora_fim
        FROM cardapios c
        LEFT JOIN cardapio_refeicoes cr ON cr.cardapio_id = c.id
        WHERE cr.cardapio_id IS NULL AND c.hora_inicio != '' AND c.hora_fim != ''
    ''').fetchall()
    for ca in cardapios_antigos:
        refs_compativeis = c.execute('''
            SELECT id FROM refeicoes
            WHERE unidade_id = ? AND NOT (hora_fim <= ? OR hora_inicio >= ?)
        ''', (ca['unidade_id'], ca['hora_inicio'], ca['hora_fim'])).fetchall()
        for r_comp in refs_compativeis:
            c.execute('INSERT OR IGNORE INTO cardapio_refeicoes (cardapio_id, refeicao_id) VALUES (?, ?)', (ca['id'], r_comp['id']))

    c.execute('''
        CREATE TABLE IF NOT EXISTS registros_refeicao (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            colaborador_id INTEGER NOT NULL,
            refeicao_id INTEGER NOT NULL,
            unidade_id INTEGER NOT NULL,
            data_registro TEXT NOT NULL,
            hora_registro TEXT NOT NULL,
            custo_registrado REAL NOT NULL,
            metodo TEXT NOT NULL,
            FOREIGN KEY (colaborador_id) REFERENCES colaboradores(id),
            FOREIGN KEY (refeicao_id) REFERENCES refeicoes(id),
            FOREIGN KEY (unidade_id) REFERENCES unidades(id)
        )
    ''')

    c.execute('''
        DELETE FROM registros_refeicao
        WHERE id NOT IN (
            SELECT MIN(id)
            FROM registros_refeicao
            GROUP BY colaborador_id, refeicao_id, data_registro
        )
    ''')
    c.execute('''
        CREATE UNIQUE INDEX IF NOT EXISTS idx_unico_colab_refeicao_dia
        ON registros_refeicao (colaborador_id, refeicao_id, data_registro)
    ''')

    c.execute('''
        CREATE TABLE IF NOT EXISTS logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            data_hora TEXT NOT NULL,
            usuario_id INTEGER,
            usuario_nome TEXT,
            unidade_id INTEGER,
            tabela TEXT NOT NULL,
            acao TEXT NOT NULL,
            descricao TEXT NOT NULL
        )
    ''')

    unid = c.execute('SELECT id FROM unidades LIMIT 1').fetchone()
    if not unid:
        c.execute("INSERT INTO unidades (descricao, ativo) VALUES ('Unidade Principal - Sede', 1)")
        unidade_id = c.lastrowid
        c.execute("INSERT INTO configuracoes (unidade_id, tema, logo_path) VALUES (?, 'tema_1', '')", (unidade_id,))
    else:
        unidade_id = unid['id']

    usr = c.execute('SELECT id, permissoes FROM usuarios ORDER BY id ASC LIMIT 1').fetchone()
    if not usr:
        todas_perms = json.dumps([
            'multiempresa', 'dashboard', 'importacao', 'cad_usuario', 'cad_refeicao',
            'cad_grupo', 'cad_colaborador', 'cad_cardapio', 'cad_unidade',
            'relatorios', 'logs', 'configuracoes'
        ])
        senha_padrao = generate_password_hash('123')
        c.execute('''
            INSERT INTO usuarios (nome, cpf, cargo, senha_hash, primeiro_acesso, ativo, permissoes)
            VALUES ('Administrador Geral', '00000000000', 'Supervisor de TI', ?, 1, 1, ?)
        ''', (senha_padrao, todas_perms))
        usuario_id = c.lastrowid
        c.execute('INSERT INTO usuario_unidades (usuario_id, unidade_id) VALUES (?, ?)', (usuario_id, unidade_id))
    else:
        # Garante que o primeiro Administrador receba as novas permissões (multiempresa, dashboard, importacao)
        perms_admin = json.loads(usr['permissoes'] or '[]')
        alterou_admin = False
        for nova_p in ('multiempresa', 'dashboard', 'importacao'):
            if nova_p not in perms_admin:
                perms_admin.append(nova_p)
                alterou_admin = True
        if alterou_admin:
            c.execute('UPDATE usuarios SET permissoes = ? WHERE id = ?', (json.dumps(perms_admin), usr['id']))

    conn.commit()

    colabs_sem_vetor = c.execute(
        "SELECT id, foto_path FROM colaboradores WHERE foto_path != '' AND (face_embeddings IS NULL OR face_embeddings = '[]')"
    ).fetchall()
    for col in colabs_sem_vetor:
        caminho_img = os.path.join(BASE_DIR, 'static', col['foto_path'])
        if os.path.exists(caminho_img):
            img_bgr = cv2.imread(caminho_img)
            vetor, _, _ = analisar_rosto_bgr(img_bgr)
            if vetor:
                c.execute("UPDATE colaboradores SET face_embeddings = ? WHERE id = ?", (json.dumps([vetor]), col['id']))
    conn.commit()
    conn.close()

init_db()
obter_motores_faciais()

def login_obrigatorio(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'usuario_id' not in session or 'unidade_id' not in session:
            return redirect(url_for('login'))
        if session.get('primeiro_acesso') == 1 and request.endpoint not in ('trocar_senha', 'logout', 'api_trocar_senha'):
            return redirect(url_for('trocar_senha'))
        return f(*args, **kwargs)
    return decorated_function

def requer_permissao(perm):
    def decorator(f):
        @wraps(f)
        def decorated_function(*args, **kwargs):
            perms = session.get('permissoes', [])
            if perm not in perms:
                return jsonify({'status': 'erro', 'mensagem': 'Você não possui permissão para realizar esta operação.'}), 403
            return f(*args, **kwargs)
        return decorated_function
    return decorator

def obter_config_unidade(unidade_id):
    conn = get_db()
    cfg = conn.execute('SELECT * FROM configuracoes WHERE unidade_id = ?', (unidade_id,)).fetchone()
    unid = conn.execute('SELECT * FROM unidades WHERE id = ?', (unidade_id,)).fetchone()
    conn.close()
    tema_chave = cfg['tema'] if cfg and cfg['tema'] in TEMAS else 'tema_1'
    logo_path = cfg['logo_path'] if cfg and cfg['logo_path'] else ''
    return {
        'tema_chave': tema_chave,
        'tema': TEMAS[tema_chave],
        'logo_path': logo_path,
        'unidade_nome': unid['descricao'] if unid else 'Unidade'
    }

@app.context_processor
def injetar_globais():
    uid = session.get('unidade_id') or session.get('totem_unidade_id') or 1
    cfg = obter_config_unidade(uid)
    return {
        'cfg_global': cfg,
        'temas_disponiveis': TEMAS,
        'csrf_token': session.get('csrf_token', '')
    }

# ==============================================================================
# ROTAS DE LOGIN, BUSCA DE UNIDADES POR CPF E REDEFINIÇÃO DE SENHA
# ==============================================================================
@app.route('/')
def index():
    if 'usuario_id' in session:
        return redirect(url_for('admin_painel'))
    return redirect(url_for('login'))

@app.route('/login', methods=['GET'])
def login():
    return render_template('login.html')

@app.route('/api/unidades-por-cpf', methods=['POST'])
def api_unidades_por_cpf():
    cpf = limpar_cpf_cnpj(request.form.get('cpf'))
    if len(cpf) != 11:
        return jsonify({'status': 'sucesso', 'unidades': []})

    conn = get_db()
    user = conn.execute('SELECT id, ativo, permissoes FROM usuarios WHERE cpf = ?', (cpf,)).fetchone()
    if not user or user['ativo'] == 0:
        conn.close()
        return jsonify({'status': 'sucesso', 'unidades': []})

    perms = json.loads(user['permissoes'] or '[]')
    if 'multiempresa' in perms:
        unidades = conn.execute('SELECT id, descricao FROM unidades WHERE ativo = 1 ORDER BY descricao').fetchall()
    else:
        unidades = conn.execute('''
            SELECT u.id, u.descricao
            FROM unidades u
            JOIN usuario_unidades uu ON uu.unidade_id = u.id
            WHERE uu.usuario_id = ? AND u.ativo = 1
            ORDER BY u.descricao
        ''', (user['id'],)).fetchall()
    conn.close()

    return jsonify({'status': 'sucesso', 'unidades': [dict(u) for u in unidades]})

@app.route('/api/login', methods=['POST'])
def api_login():
    cpf = limpar_cpf_cnpj(request.form.get('cpf'))
    senha = request.form.get('senha', '')
    unidade_id = request.form.get('unidade_id', type=int)

    if not cpf or not senha or not unidade_id:
        return jsonify({'status': 'erro', 'mensagem': 'Preencha CPF, senha e selecione a unidade.'})

    conn = get_db()
    user = conn.execute('SELECT * FROM usuarios WHERE cpf = ?', (cpf,)).fetchone()
    if not user or not check_password_hash(user['senha_hash'], senha):
        conn.close()
        return jsonify({'status': 'erro', 'mensagem': 'CPF ou senha incorretos.'})

    if user['ativo'] == 0:
        conn.close()
        return jsonify({'status': 'erro', 'mensagem': 'Este usuário está inativo. Contate o administrador.'})

    perms = json.loads(user['permissoes'] or '[]')
    eh_multi = 'multiempresa' in perms

    unid = conn.execute('SELECT * FROM unidades WHERE id = ? AND ativo = 1', (unidade_id,)).fetchone()
    if not unid:
        conn.close()
        return jsonify({'status': 'erro', 'mensagem': 'Unidade inválida ou inativa.'})

    if not eh_multi:
        vinculo = conn.execute(
            'SELECT 1 FROM usuario_unidades WHERE usuario_id = ? AND unidade_id = ?',
            (user['id'], unidade_id)
        ).fetchone()
        if not vinculo:
            conn.close()
            return jsonify({'status': 'erro', 'mensagem': 'Você não possui permissão de acesso para a unidade selecionada.'})

    conn.close()

    session['usuario_id'] = user['id']
    session['usuario_nome'] = user['nome']
    session['unidade_id'] = unidade_id
    session['unidade_nome'] = unid['descricao']
    session['primeiro_acesso'] = user['primeiro_acesso']
    session['permissoes'] = perms

    registrar_log('usuarios', 'LOGIN', f"Usuário '{user['nome']}' acessou a unidade '{unid['descricao']}'.", unidade_id)

    if user['primeiro_acesso'] == 1:
        return jsonify({'status': 'sucesso', 'redirect': url_for('trocar_senha')})

    aba_inicial = obter_aba_inicial_permitida(perms)
    return jsonify({'status': 'sucesso', 'redirect': url_for('admin_painel', aba=aba_inicial)})

@app.route('/trocar-senha', methods=['GET'])
@login_obrigatorio
def trocar_senha():
    return render_template('trocar_senha.html')

@app.route('/api/trocar-senha', methods=['POST'])
@login_obrigatorio
def api_trocar_senha():
    nova_senha = request.form.get('nova_senha', '')
    confirmar = request.form.get('confirmar_senha', '')

    if nova_senha != confirmar:
        return jsonify({'status': 'erro', 'mensagem': 'A confirmação de senha não confere.'})

    if len(nova_senha) < 8 or not re.search(r'[A-Za-z]', nova_senha) or not re.search(r'\d', nova_senha) or not re.search(r'[^A-Za-z0-9]', nova_senha):
        return jsonify({'status': 'erro', 'mensagem': 'A nova senha deve conter no mínimo 8 caracteres, incluindo letras, números e símbolos.'})

    if nova_senha == '123':
        return jsonify({'status': 'erro', 'mensagem': 'A nova senha não pode ser igual à senha padrão.'})

    novo_hash = generate_password_hash(nova_senha)
    conn = get_db()
    conn.execute('UPDATE usuarios SET senha_hash = ?, primeiro_acesso = 0 WHERE id = ?', (novo_hash, session['usuario_id']))
    conn.commit()
    conn.close()

    session['primeiro_acesso'] = 0
    registrar_log('usuarios', 'ALTERAÇÃO', f"Usuário '{session['usuario_nome']}' redefiniu sua senha obrigatória de primeiro acesso.")
    aba_inicial = obter_aba_inicial_permitida(session.get('permissoes', []))
    return jsonify({'status': 'sucesso', 'mensagem': 'Senha atualizada com sucesso!', 'redirect': url_for('admin_painel', aba=aba_inicial)})

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

# ==============================================================================
# MÓDULO 1: PAINEL ADMINISTRATIVO
# ==============================================================================
@app.route('/admin')
@login_obrigatorio
def admin_painel():
    uid = session['unidade_id']
    conn = get_db()

    # Sincroniza as permissões atuais do usuário logado diretamente do banco
    usr_atual = conn.execute('SELECT permissoes, ativo FROM usuarios WHERE id = ?', (session['usuario_id'],)).fetchone()
    if not usr_atual or usr_atual['ativo'] == 0:
        conn.close()
        session.clear()
        return redirect(url_for('login'))

    perms = json.loads(usr_atual['permissoes'] or '[]')
    session['permissoes'] = perms
    eh_multi = 'multiempresa' in perms

    aba_solicitada = request.args.get('aba', '')
    mapa_dict = dict(MAPA_ABAS_PERMISSOES)

    # Verifica se o usuário tem permissão para a aba solicitada (ex: dashboard); se não tiver, redireciona para a primeira permitida
    if not aba_solicitada or (aba_solicitada in mapa_dict and mapa_dict[aba_solicitada] not in perms):
        aba = obter_aba_inicial_permitida(perms)
    else:
        aba = aba_solicitada

    if eh_multi:
        unidades = conn.execute('SELECT * FROM unidades ORDER BY descricao').fetchall()
        unidades_ativas = conn.execute('SELECT * FROM unidades WHERE ativo = 1 ORDER BY descricao').fetchall()
    else:
        unidades = conn.execute('SELECT * FROM unidades WHERE id = ?', (uid,)).fetchall()
        unidades_ativas = unidades

    grupos = conn.execute('SELECT * FROM grupos WHERE unidade_id = ? ORDER BY descricao', (uid,)).fetchall()
    refeicoes = conn.execute('SELECT * FROM refeicoes WHERE unidade_id = ? ORDER BY hora_inicio', (uid,)).fetchall()

    vinculos_rg = conn.execute('''
        SELECT rg.refeicao_id, rg.grupo_id, g.descricao as grupo_nome, r.descricao as refeicao_nome
        FROM refeicao_grupos rg
        JOIN grupos g ON g.id = rg.grupo_id
        JOIN refeicoes r ON r.id = rg.refeicao_id
        WHERE g.unidade_id = ? AND r.unidade_id = ?
    ''', (uid, uid)).fetchall()

    mapa_ref_grupos = {}
    mapa_grupo_refs = {}
    for v in vinculos_rg:
        mapa_ref_grupos.setdefault(v['refeicao_id'], []).append({'id': v['grupo_id'], 'nome': v['grupo_nome']})
        mapa_grupo_refs.setdefault(v['grupo_id'], []).append({'id': v['refeicao_id'], 'nome': v['refeicao_nome']})

    vinculos_cr = conn.execute('''
        SELECT cr.cardapio_id, r.id as refeicao_id, r.descricao as refeicao_nome, r.hora_inicio, r.hora_fim
        FROM cardapio_refeicoes cr
        JOIN refeicoes r ON r.id = cr.refeicao_id
        JOIN cardapios c ON c.id = cr.cardapio_id
        WHERE c.unidade_id = ? AND r.unidade_id = ?
        ORDER BY r.hora_inicio ASC
    ''', (uid, uid)).fetchall()

    mapa_cardapio_refs = {}
    for vc in vinculos_cr:
        mapa_cardapio_refs.setdefault(vc['cardapio_id'], []).append({
            'id': vc['refeicao_id'],
            'nome': vc['refeicao_nome'],
            'hora_inicio': vc['hora_inicio'],
            'hora_fim': vc['hora_fim']
        })

    colaboradores_db = conn.execute('''
        SELECT c.*, g.descricao as grupo_nome
        FROM colaboradores c
        LEFT JOIN grupos g ON g.id = c.grupo_id
        WHERE c.unidade_id = ?
        ORDER BY c.nome
    ''', (uid,)).fetchall()

    colaboradores = []
    for c in colaboradores_db:
        qtd_vetores = len(json.loads(c['face_embeddings'] or '[]'))
        d_col = dict(c)
        d_col['qtd_vetores'] = qtd_vetores
        colaboradores.append(d_col)

    cardapios = conn.execute('''
        SELECT * FROM cardapios
        WHERE unidade_id = ?
        ORDER BY data_vigencia DESC, id DESC
    ''', (uid,)).fetchall()

    if eh_multi:
        usuarios_db = conn.execute('''
            SELECT DISTINCT u.* FROM usuarios u
            LEFT JOIN usuario_unidades uu ON uu.usuario_id = u.id
            WHERE uu.unidade_id = ? OR u.permissoes LIKE '%"multiempresa"%'
            ORDER BY u.nome
        ''', (uid,)).fetchall()
    else:
        usuarios_db = conn.execute('''
            SELECT u.* FROM usuarios u
            JOIN usuario_unidades uu ON uu.usuario_id = u.id
            WHERE uu.unidade_id = ? AND u.permissoes NOT LIKE '%"multiempresa"%'
            ORDER BY u.nome
        ''', (uid,)).fetchall()

    usuarios = []
    for u in usuarios_db:
        if eh_multi:
            unids_u = conn.execute('''
                SELECT u.id, u.descricao FROM usuario_unidades uu
                JOIN unidades u ON u.id = uu.unidade_id
                WHERE uu.usuario_id = ?
            ''', (u['id'],)).fetchall()
        else:
            unids_u = conn.execute('''
                SELECT u.id, u.descricao FROM usuario_unidades uu
                JOIN unidades u ON u.id = uu.unidade_id
                WHERE uu.usuario_id = ? AND uu.unidade_id = ?
            ''', (u['id'], uid)).fetchall()

        usuarios.append({
            'id': u['id'],
            'nome': u['nome'],
            'cpf': u['cpf'],
            'cargo': u['cargo'],
            'ativo': u['ativo'],
            'permissoes': json.loads(u['permissoes']),
            'unidades': [dict(x) for x in unids_u],
            'unidades_ids': [x['id'] for x in unids_u]
        })

    hoje = datetime.now().strftime('%Y-%m-%d')
    resumo_hoje = conn.execute('''
        SELECT COUNT(*) as total_refeicoes, COALESCE(SUM(custo_registrado), 0) as custo_total
        FROM registros_refeicao
        WHERE unidade_id = ? AND data_registro = ?
    ''', (uid, hoje)).fetchone()

    conn.close()
    return render_template(
        'admin.html',
        aba=aba,
        eh_multi=eh_multi,
        unidades=unidades,
        unidades_ativas=unidades_ativas,
        grupos=grupos,
        refeicoes=refeicoes,
        mapa_ref_grupos=mapa_ref_grupos,
        mapa_grupo_refs=mapa_grupo_refs,
        mapa_cardapio_refs=mapa_cardapio_refs,
        colaboradores=colaboradores,
        cardapios=cardapios,
        usuarios=usuarios,
        resumo_hoje=resumo_hoje,
        hoje=hoje
    )

# ==============================================================================
# IMPORTAÇÃO E EXPORTAÇÃO DE PLANILHAS EXCEL (.XLSX)
# ==============================================================================
def interpretar_ativo_excel(valor):
    s = str(valor or 'Sim').strip().lower()
    if s in ('nao', 'não', '0', 'inativo', 'false', 'n'):
        return 0
    return 1

def normalizar_hora_excel(valor):
    if valor is None:
        return ''
    if isinstance(valor, datetime):
        return valor.strftime('%H:%M')
    if isinstance(valor, time):
        return valor.strftime('%H:%M')
    s = str(valor).strip()
    m = re.match(r'^(\d{1,2}):(\d{2})', s)
    if m:
        hh, mm = int(m.group(1)), int(m.group(2))
        if 0 <= hh <= 23 and 0 <= mm <= 59:
            return f"{hh:02d}:{mm:02d}"
    return ''

def normalizar_data_excel(valor):
    if valor is None:
        return ''
    if isinstance(valor, datetime):
        return valor.strftime('%Y-%m-%d')
    if isinstance(valor, date):
        return valor.strftime('%Y-%m-%d')
    s = str(valor).strip()
    if re.match(r'^\d{4}-\d{2}-\d{2}$', s):
        return s
    m_br = re.match(r'^(\d{2})/(\d{2})/(\d{4})$', s)
    if m_br:
        return f"{m_br.group(3)}-{m_br.group(2)}-{m_br.group(1)}"
    return ''

@app.route('/admin/importacao/modelo/<tipo>', methods=['GET'])
@login_obrigatorio
@requer_permissao('importacao')
def baixar_modelo_excel(tipo):
    uid = session['unidade_id']
    eh_multi = usuario_eh_multiempresa()

    if tipo == 'unidades' and not eh_multi:
        return "Apenas usuários MultiEmpresa podem baixar o modelo de unidades.", 403

    wb = Workbook()
    ws = wb.active
    ws.title = "Dados_Importacao"

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    ex_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    conn = get_db()
    grupos_unid = [g['descricao'] for g in conn.execute('SELECT descricao FROM grupos WHERE unidade_id = ? ORDER BY descricao', (uid,)).fetchall()]
    refs_unid = [r['descricao'] for r in conn.execute('SELECT descricao FROM refeicoes WHERE unidade_id = ? ORDER BY hora_inicio', (uid,)).fetchall()]
    unids_todas = [u['descricao'] for u in conn.execute('SELECT descricao FROM unidades WHERE ativo = 1 ORDER BY descricao').fetchall()]
    conn.close()

    if tipo == 'usuarios':
        colunas = ['Nome', 'CPF', 'Cargo', 'Ativo (Sim/Nao)', 'Permissoes (Separadas por virgula)', 'Unidades (Separadas por virgula)']
        exemplo = ['Maria Souza (EXEMPLO - APAGUE ESTA LINHA)', '12345678901', 'Nutricionista', 'Sim', 'dashboard, cad_cardapio, relatorios', session.get('unidade_nome', 'Unidade Principal - Sede')]
    elif tipo == 'refeicoes':
        colunas = ['Descricao', 'Hora_Inicio (HH:MM)', 'Hora_Fim (HH:MM)', 'Custo (R$)', 'Grupos_Vinculados (Separados por virgula)', 'Ativo (Sim/Nao)']
        grp_ex = grupos_unid[0] if grupos_unid else 'Administrativo'
        exemplo = ['Almoço (EXEMPLO - APAGUE ESTA LINHA)', '11:00', '14:00', '16.50', grp_ex, 'Sim']
    elif tipo == 'grupos':
        colunas = ['Descricao', 'Refeicoes_Vinculadas (Separadas por virgula)', 'Ativo (Sim/Nao)']
        ref_ex = ", ".join(refs_unid[:2]) if refs_unid else 'Almoço, Jantar'
        exemplo = ['Plantonistas Diurnos (EXEMPLO - APAGUE ESTA LINHA)', ref_ex, 'Sim']
    elif tipo == 'colaboradores':
        colunas = ['Nome', 'CPF', 'Grupo', 'Ativo (Sim/Nao)']
        grp_ex = grupos_unid[0] if grupos_unid else 'Administrativo'
        exemplo = ['João da Silva (EXEMPLO - APAGUE ESTA LINHA)', '98765432100', grp_ex, 'Sim']
    elif tipo == 'cardapios':
        colunas = ['Data_Vigencia (DD/MM/AAAA)', 'Refeicoes_Vinculadas (Separadas por virgula)', 'Proteinas', 'Guarnicao', 'Saladas', 'Sobremesas', 'Bebidas', 'Ativo (Sim/Nao)']
        ref_ex = refs_unid[0] if refs_unid else 'Almoço'
        exemplo = [datetime.now().strftime('%d/%m/%Y'), ref_ex, 'Frango Grelhado / Iscas Aceboladas', 'Arroz, Feijão e Purê', 'Alface, Tomate e Cenoura', 'Fruta da Estação', 'Suco Natural', 'Sim']
    elif tipo == 'unidades':
        colunas = ['Descricao', 'Ativo (Sim/Nao)']
        exemplo = ['Nova Unidade Filial (EXEMPLO - APAGUE ESTA LINHA)', 'Sim']
    else:
        return "Tipo de planilha inválido.", 400

    ws.append(colunas)
    ws.append(exemplo)

    for col_num in range(1, len(colunas) + 1):
        cell_h = ws.cell(row=1, column=col_num)
        cell_h.fill = header_fill
        cell_h.font = header_font
        cell_h.alignment = Alignment(horizontal="center", vertical="center")

        cell_ex = ws.cell(row=2, column=col_num)
        cell_ex.fill = ex_fill

        largura = max(len(str(colunas[col_num - 1])), len(str(exemplo[col_num - 1]))) + 6
        ws.column_dimensions[cell_h.column_letter].width = min(largura, 50)

    # 2ª Aba com Instruções e Referências da Unidade Atual
    ws_info = wb.create_sheet(title="Instrucoes_e_Referencias")
    ws_info.append(["INSTRUÇÕES PARA IMPORTAÇÃO SEGURA"])
    ws_info.cell(row=1, column=1).font = Font(bold=True, size=12)
    ws_info.append(["1. Preencha os dados na primeira aba (Dados_Importacao) a partir da linha 2 (substitua ou apague a linha de exemplo)."])
    ws_info.append(["2. Cadastros já existentes no sistema NÃO serão alterados e registros duplicados serão ignorados automaticamente."])
    ws_info.append(["3. Unidades de destino dos cadastros: " + session.get('unidade_nome', '')])
    ws_info.append([])
    ws_info.append(["GRUPOS CADASTRADOS NESTA UNIDADE:"] + (grupos_unid if grupos_unid else ["Nenhum grupo cadastrado ainda"]))
    ws_info.append(["REFEIÇÕES CADASTRADAS NESTA UNIDADE:"] + (refs_unid if refs_unid else ["Nenhuma refeição cadastrada ainda"]))
    if tipo == 'usuarios':
        ws_info.append(["CÓDIGOS DE PERMISSÕES VÁLIDOS:", "dashboard, cad_usuario, cad_refeicao, cad_grupo, cad_colaborador, cad_cardapio, cad_unidade, importacao, relatorios, logs, configuracoes" + (", multiempresa" if eh_multi else "")])
        if eh_multi:
            ws_info.append(["UNIDADES ATIVAS DISPONÍVEIS:"] + unids_todas)
    ws_info.column_dimensions['A'].width = 90

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name=f"modelo_importacao_{tipo}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@app.route('/api/importacao/processar/<tipo>', methods=['POST'])
@login_obrigatorio
@requer_permissao('importacao')
def processar_importacao_excel(tipo):
    uid = session['unidade_id']
    eh_multi = usuario_eh_multiempresa()

    if tipo == 'unidades' and not eh_multi:
        return jsonify({'status': 'erro', 'mensagem': 'Apenas usuários com a permissão MultiEmpresa podem importar novas unidades.'}), 403

    arquivo = request.files.get('arquivo_excel')
    if not arquivo or not arquivo.filename.lower().endswith('.xlsx'):
        return jsonify({'status': 'erro', 'mensagem': 'Selecione um arquivo Excel válido no formato .xlsx.'})

    try:
        wb = load_workbook(arquivo, data_only=True)
        ws = wb.active
    except Exception:
        return jsonify({'status': 'erro', 'mensagem': 'Não foi possível ler a planilha. Certifique-se de enviar um arquivo .xlsx válido.'})

    linhas = list(ws.iter_rows(values_only=True))
    if len(linhas) <= 1:
        return jsonify({'status': 'erro', 'mensagem': 'A planilha enviada está vazia (sem linhas de dados abaixo do cabeçalho).'})

    inseridos = 0
    ignorados = 0
    detalhes = []

    conn = get_db()

    try:
        # ------------------------------------------------------------------
        # 1. IMPORTAÇÃO DE UNIDADES
        # ------------------------------------------------------------------
        if tipo == 'unidades':
            existentes = {
                r['descricao'].strip().lower()
                for r in conn.execute('SELECT descricao FROM unidades').fetchall()
            }
            vistos_planilha = set()

            for idx, row in enumerate(linhas[1:], start=2):
                if not row or all(cel is None or str(cel).strip() == '' for cel in row):
                    continue
                descricao = str(row[0] or '').strip()
                if not descricao or 'EXEMPLO - APAGUE ESTA LINHA' in descricao.upper():
                    continue
                ativo = interpretar_ativo_excel(row[1] if len(row) > 1 else 'Sim')
                chave = descricao.lower()

                if chave in existentes or chave in vistos_planilha:
                    ignorados += 1
                    detalhes.append(f"Linha {idx}: Unidade '{descricao}' ignorada (já cadastrada ou duplicada na planilha).")
                    continue

                cur = conn.execute('INSERT INTO unidades (descricao, ativo) VALUES (?, ?)', (descricao, ativo))
                novo_id = cur.lastrowid
                conn.execute("INSERT OR IGNORE INTO configuracoes (unidade_id, tema, logo_path) VALUES (?, 'tema_1', '')", (novo_id,))
                conn.execute('INSERT OR IGNORE INTO usuario_unidades (usuario_id, unidade_id) VALUES (?, ?)', (session['usuario_id'], novo_id))
                vistos_planilha.add(chave)
                inseridos += 1

        # ------------------------------------------------------------------
        # 2. IMPORTAÇÃO DE GRUPOS
        # ------------------------------------------------------------------
        elif tipo == 'grupos':
            existentes = {
                r['descricao'].strip().lower()
                for r in conn.execute('SELECT descricao FROM grupos WHERE unidade_id = ?', (uid,)).fetchall()
            }
            mapa_refs = {
                r['descricao'].strip().lower(): r['id']
                for r in conn.execute('SELECT id, descricao FROM refeicoes WHERE unidade_id = ?', (uid,)).fetchall()
            }
            vistos_planilha = set()

            for idx, row in enumerate(linhas[1:], start=2):
                if not row or all(cel is None or str(cel).strip() == '' for cel in row):
                    continue
                descricao = str(row[0] or '').strip()
                if not descricao or 'EXEMPLO - APAGUE ESTA LINHA' in descricao.upper():
                    continue
                refs_str = str(row[1] or '').strip() if len(row) > 1 else ''
                ativo = interpretar_ativo_excel(row[2] if len(row) > 2 else 'Sim')
                chave = descricao.lower()

                if chave in existentes or chave in vistos_planilha:
                    ignorados += 1
                    detalhes.append(f"Linha {idx}: Grupo '{descricao}' ignorado (já existente na unidade ou duplicado na planilha).")
                    continue

                cur = conn.execute('INSERT INTO grupos (descricao, ativo, unidade_id) VALUES (?, ?, ?)', (descricao, ativo, uid))
                novo_grp_id = cur.lastrowid

                if refs_str:
                    for nome_ref in [x.strip().lower() for x in refs_str.split(',') if x.strip()]:
                        if nome_ref in mapa_refs:
                            conn.execute('INSERT OR IGNORE INTO refeicao_grupos (refeicao_id, grupo_id) VALUES (?, ?)', (mapa_refs[nome_ref], novo_grp_id))

                vistos_planilha.add(chave)
                inseridos += 1

        # ------------------------------------------------------------------
        # 3. IMPORTAÇÃO DE REFEIÇÕES
        # ------------------------------------------------------------------
        elif tipo == 'refeicoes':
            existentes = {
                r['descricao'].strip().lower()
                for r in conn.execute('SELECT descricao FROM refeicoes WHERE unidade_id = ?', (uid,)).fetchall()
            }
            mapa_grupos = {
                g['descricao'].strip().lower(): g['id']
                for g in conn.execute('SELECT id, descricao FROM grupos WHERE unidade_id = ?', (uid,)).fetchall()
            }
            vistos_planilha = set()

            for idx, row in enumerate(linhas[1:], start=2):
                if not row or all(cel is None or str(cel).strip() == '' for cel in row):
                    continue
                descricao = str(row[0] or '').strip()
                if not descricao or 'EXEMPLO - APAGUE ESTA LINHA' in descricao.upper():
                    continue

                hora_ini = normalizar_hora_excel(row[1] if len(row) > 1 else None)
                hora_fim = normalizar_hora_excel(row[2] if len(row) > 2 else None)
                custo_raw = str(row[3] if len(row) > 3 and row[3] is not None else '0').replace('R$', '').replace(',', '.').strip()
                grupos_str = str(row[4] or '').strip() if len(row) > 4 else ''
                ativo = interpretar_ativo_excel(row[5] if len(row) > 5 else 'Sim')

                if not hora_ini or not hora_fim:
                    ignorados += 1
                    detalhes.append(f"Linha {idx}: Refeição '{descricao}' ignorada (horário inicial ou final inválido; use HH:MM).")
                    continue

                try:
                    custo_val = float(custo_raw)
                except ValueError:
                    ignorados += 1
                    detalhes.append(f"Linha {idx}: Refeição '{descricao}' ignorada (custo inválido).")
                    continue

                chave = descricao.lower()
                if chave in existentes or chave in vistos_planilha:
                    ignorados += 1
                    detalhes.append(f"Linha {idx}: Refeição '{descricao}' ignorada (já cadastrada na unidade ou duplicada na planilha).")
                    continue

                cur = conn.execute('''
                    INSERT INTO refeicoes (descricao, hora_inicio, hora_fim, custo, ativo, unidade_id)
                    VALUES (?, ?, ?, ?, ?, ?)
                ''', (descricao, hora_ini, hora_fim, custo_val, ativo, uid))
                novo_ref_id = cur.lastrowid

                if grupos_str:
                    for nome_grp in [x.strip().lower() for x in grupos_str.split(',') if x.strip()]:
                        if nome_grp in mapa_grupos:
                            conn.execute('INSERT OR IGNORE INTO refeicao_grupos (refeicao_id, grupo_id) VALUES (?, ?)', (novo_ref_id, mapa_grupos[nome_grp]))

                vistos_planilha.add(chave)
                inseridos += 1

        # ------------------------------------------------------------------
        # 4. IMPORTAÇÃO DE COLABORADORES
        # ------------------------------------------------------------------
        elif tipo == 'colaboradores':
            existentes = {
                r['cpf'].strip()
                for r in conn.execute('SELECT cpf FROM colaboradores WHERE unidade_id = ?', (uid,)).fetchall()
            }
            mapa_grupos = {
                g['descricao'].strip().lower(): g['id']
                for g in conn.execute('SELECT id, descricao FROM grupos WHERE unidade_id = ?', (uid,)).fetchall()
            }
            vistos_planilha = set()

            for idx, row in enumerate(linhas[1:], start=2):
                if not row or all(cel is None or str(cel).strip() == '' for cel in row):
                    continue
                nome = str(row[0] or '').strip()
                if not nome or 'EXEMPLO - APAGUE ESTA LINHA' in nome.upper():
                    continue

                cpf = limpar_cpf_cnpj(row[1] if len(row) > 1 else '')
                if len(cpf) > 0 and len(cpf) < 11:
                    cpf = cpf.zfill(11)

                nome_grp = str(row[2] or '').strip().lower() if len(row) > 2 else ''
                ativo = interpretar_ativo_excel(row[3] if len(row) > 3 else 'Sim')

                if len(cpf) != 11:
                    ignorados += 1
                    detalhes.append(f"Linha {idx}: Colaborador '{nome}' ignorado (CPF inválido).")
                    continue

                if cpf in existentes or cpf in vistos_planilha:
                    ignorados += 1
                    detalhes.append(f"Linha {idx}: Colaborador '{nome}' (CPF {cpf}) ignorado (CPF já cadastrado nesta unidade ou duplicado na planilha).")
                    continue

                if nome_grp not in mapa_grupos:
                    ignorados += 1
                    detalhes.append(f"Linha {idx}: Colaborador '{nome}' ignorado (Grupo '{row[2] if len(row) > 2 else ''}' não encontrado nesta unidade).")
                    continue

                conn.execute('''
                    INSERT INTO colaboradores (nome, cpf, foto_path, face_embeddings, grupo_id, ativo, unidade_id)
                    VALUES (?, ?, '', '[]', ?, ?, ?)
                ''', (nome, cpf, mapa_grupos[nome_grp], ativo, uid))

                vistos_planilha.add(cpf)
                inseridos += 1

        # ------------------------------------------------------------------
        # 5. IMPORTAÇÃO DE CARDÁPIOS
        # ------------------------------------------------------------------
        elif tipo == 'cardapios':
            mapa_refs = {
                r['descricao'].strip().lower(): dict(r)
                for r in conn.execute('SELECT id, descricao, hora_inicio, hora_fim FROM refeicoes WHERE unidade_id = ?', (uid,)).fetchall()
            }
            # Pares (data_vigencia, refeicao_id) já ocupados na unidade
            ocupados_db = set()
            rows_ocup = conn.execute('''
                SELECT c.data_vigencia, cr.refeicao_id
                FROM cardapios c
                JOIN cardapio_refeicoes cr ON cr.cardapio_id = c.id
                WHERE c.unidade_id = ?
            ''', (uid,)).fetchall()
            for r_oc in rows_ocup:
                ocupados_db.add((r_oc['data_vigencia'], r_oc['refeicao_id']))

            for idx, row in enumerate(linhas[1:], start=2):
                if not row or all(cel is None or str(cel).strip() == '' for cel in row):
                    continue
                proteinas = str(row[2] or '').strip() if len(row) > 2 else ''
                if 'EXEMPLO - APAGUE ESTA LINHA' in proteinas.upper():
                    continue

                data_vig = normalizar_data_excel(row[0] if len(row) > 0 else None)
                refs_str = str(row[1] or '').strip() if len(row) > 1 else ''
                guarnicao = str(row[3] or '').strip() if len(row) > 3 else ''
                saladas = str(row[4] or '').strip() if len(row) > 4 else ''
                sobremesas = str(row[5] or '').strip() if len(row) > 5 else ''
                bebidas = str(row[6] or '').strip() if len(row) > 6 else ''
                ativo = interpretar_ativo_excel(row[7] if len(row) > 7 else 'Sim')

                if not data_vig or not all([refs_str, proteinas, guarnicao, saladas, sobremesas, bebidas]):
                    ignorados += 1
                    detalhes.append(f"Linha {idx}: Cardápio ignorado (data inválida ou campos obrigatórios em branco).")
                    continue

                nomes_refs = [x.strip().lower() for x in refs_str.split(',') if x.strip()]
                objs_ref = [mapa_refs[n] for n in nomes_refs if n in mapa_refs]
                if not objs_ref or len(objs_ref) != len(nomes_refs):
                    ignorados += 1
                    detalhes.append(f"Linha {idx}: Cardápio de {data_vig} ignorado (uma ou mais refeições informadas não existem na unidade).")
                    continue

                # Verifica se alguma refeição nessa data já possui cardápio cadastrado (não modifica antigos)
                conflito = any((data_vig, o['id']) in ocupados_db for o in objs_ref)
                if conflito:
                    ignorados += 1
                    detalhes.append(f"Linha {idx}: Cardápio de {data_vig} ignorado (já existe cardápio cadastrado para esta data e refeição).")
                    continue

                objs_ref.sort(key=lambda x: x['hora_inicio'])
                cur = conn.execute('''
                    INSERT INTO cardapios (data_vigencia, hora_inicio, hora_fim, guarnicao, proteinas, saladas, sobremesas, bebidas, ativo, unidade_id)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ''', (data_vig, objs_ref[0]['hora_inicio'], objs_ref[-1]['hora_fim'], guarnicao, proteinas, saladas, sobremesas, bebidas, ativo, uid))
                novo_card_id = cur.lastrowid

                for o in objs_ref:
                    conn.execute('INSERT INTO cardapio_refeicoes (cardapio_id, refeicao_id) VALUES (?, ?)', (novo_card_id, o['id']))
                    ocupados_db.add((data_vig, o['id']))

                inseridos += 1

        # ------------------------------------------------------------------
        # 6. IMPORTAÇÃO DE USUÁRIOS
        # ------------------------------------------------------------------
        elif tipo == 'usuarios':
            existentes = {
                r['cpf'].strip()
                for r in conn.execute('SELECT cpf FROM usuarios').fetchall()
            }
            mapa_unidades = {
                u['descricao'].strip().lower(): u['id']
                for u in conn.execute('SELECT id, descricao FROM unidades').fetchall()
            }
            perms_validas = {
                'dashboard', 'cad_usuario', 'cad_refeicao', 'cad_grupo',
                'cad_colaborador', 'cad_cardapio', 'cad_unidade',
                'importacao', 'relatorios', 'logs', 'configuracoes'
            }
            if eh_multi:
                perms_validas.add('multiempresa')

            vistos_planilha = set()
            senha_padrao_hash = generate_password_hash('123')

            for idx, row in enumerate(linhas[1:], start=2):
                if not row or all(cel is None or str(cel).strip() == '' for cel in row):
                    continue
                nome = str(row[0] or '').strip()
                if not nome or 'EXEMPLO - APAGUE ESTA LINHA' in nome.upper():
                    continue

                cpf = limpar_cpf_cnpj(row[1] if len(row) > 1 else '')
                if len(cpf) > 0 and len(cpf) < 11:
                    cpf = cpf.zfill(11)
                cargo = str(row[2] or '').strip() if len(row) > 2 else ''
                ativo = interpretar_ativo_excel(row[3] if len(row) > 3 else 'Sim')
                perms_str = str(row[4] or '').strip() if len(row) > 4 else 'dashboard'
                unids_str = str(row[5] or '').strip() if len(row) > 5 else ''

                if len(cpf) != 11:
                    ignorados += 1
                    detalhes.append(f"Linha {idx}: Usuário '{nome}' ignorado (CPF inválido).")
                    continue

                if cpf in existentes or cpf in vistos_planilha:
                    ignorados += 1
                    detalhes.append(f"Linha {idx}: Usuário '{nome}' (CPF {cpf}) ignorado (CPF já cadastrado no sistema ou duplicado na planilha).")
                    continue

                perms_lista = [p.strip().lower() for p in perms_str.split(',') if p.strip().lower() in perms_validas]
                if not perms_lista:
                    perms_lista = ['dashboard']

                if eh_multi and unids_str:
                    unids_ids = [mapa_unidades[n.strip().lower()] for n in unids_str.split(',') if n.strip().lower() in mapa_unidades]
                    if not unids_ids:
                        unids_ids = [uid]
                else:
                    unids_ids = [uid]

                cur = conn.execute('''
                    INSERT INTO usuarios (nome, cpf, cargo, senha_hash, primeiro_acesso, ativo, permissoes)
                    VALUES (?, ?, ?, ?, 1, ?, ?)
                ''', (nome, cpf, cargo, senha_padrao_hash, ativo, json.dumps(perms_lista)))
                novo_usr_id = cur.lastrowid

                for un_id in set(unids_ids):
                    conn.execute('INSERT OR IGNORE INTO usuario_unidades (usuario_id, unidade_id) VALUES (?, ?)', (novo_usr_id, un_id))

                vistos_planilha.add(cpf)
                inseridos += 1
        else:
            conn.close()
            return jsonify({'status': 'erro', 'mensagem': 'Categoria de importação desconhecida.'}), 400

        conn.commit()
    except Exception as e:
        conn.rollback()
        conn.close()
        return jsonify({'status': 'erro', 'mensagem': f"Erro durante o processamento da planilha: {str(e)}"})

    conn.close()
    registrar_log(tipo, 'IMPORTAÇÃO', f"Importação de planilha ({tipo}): {inseridos} inseridos, {ignorados} ignorados.")

    resumo_msg = f"Importação concluída! {inseridos} novo(s) cadastro(s) inserido(s) e {ignorados} linha(s) ignorada(s)."
    if detalhes:
        resumo_msg += "\n\nDetalhes das linhas ignoradas:\n• " + "\n• ".join(detalhes[:10])
        if len(detalhes) > 10:
            resumo_msg += f"\n... e mais {len(detalhes) - 10} linha(s)."

    return jsonify({'status': 'sucesso', 'mensagem': resumo_msg})

# ==============================================================================
# DEMAIS ROTAS DE CADASTRO (USUÁRIOS, UNIDADES, GRUPOS, REFEIÇÕES, COLABORADORES, CARDÁPIO)
# ==============================================================================
@app.route('/api/usuario/salvar', methods=['POST'])
@login_obrigatorio
@requer_permissao('cad_usuario')
def salvar_usuario():
    uid_edit = request.form.get('id', type=int)
    nome = request.form.get('nome', '').strip()
    cpf = limpar_cpf_cnpj(request.form.get('cpf'))
    cargo = request.form.get('cargo', '').strip()
    ativo = request.form.get('ativo', type=int, default=1)
    permissoes = request.form.getlist('permissoes')
    unidades_ids = request.form.getlist('unidades_ids')

    uid_logada = session['unidade_id']
    eh_multi = usuario_eh_multiempresa()

    if not eh_multi:
        unidades_ids = [str(uid_logada)]
        if 'multiempresa' in permissoes:
            permissoes.remove('multiempresa')

    if not nome or len(cpf) != 11 or not unidades_ids:
        return jsonify({'status': 'erro', 'mensagem': 'Preencha Nome, CPF válido (11 dígitos) e vincule ao menos uma Unidade.'})

    if uid_edit == session['usuario_id']:
        if 'cad_usuario' not in permissoes:
            permissoes.append('cad_usuario')
        if eh_multi and 'multiempresa' not in permissoes:
            permissoes.append('multiempresa')

    perms_json = json.dumps(permissoes)
    conn = get_db()
    try:
        if uid_edit:
            alvo = conn.execute('SELECT * FROM usuarios WHERE id = ?', (uid_edit,)).fetchone()
            if not alvo:
                conn.close()
                return jsonify({'status': 'erro', 'mensagem': 'Usuário não encontrado.'})

            if not eh_multi:
                alvo_perms = json.loads(alvo['permissoes'] or '[]')
                pertence = conn.execute(
                    'SELECT 1 FROM usuario_unidades WHERE usuario_id = ? AND unidade_id = ?',
                    (uid_edit, uid_logada)
                ).fetchone()
                if not pertence or 'multiempresa' in alvo_perms:
                    conn.close()
                    return jsonify({'status': 'erro', 'mensagem': 'Você não tem permissão para alterar este usuário.'}), 403

            conn.execute('UPDATE usuarios SET nome=?, cpf=?, cargo=?, ativo=?, permissoes=? WHERE id=?', (nome, cpf, cargo, ativo, perms_json, uid_edit))

            if eh_multi:
                conn.execute('DELETE FROM usuario_unidades WHERE usuario_id=?', (uid_edit,))
                for un_id in unidades_ids:
                    conn.execute('INSERT INTO usuario_unidades (usuario_id, unidade_id) VALUES (?, ?)', (uid_edit, int(un_id)))
            else:
                conn.execute('INSERT OR IGNORE INTO usuario_unidades (usuario_id, unidade_id) VALUES (?, ?)', (uid_edit, uid_logada))

            conn.commit()
            registrar_log('usuarios', 'ALTERAÇÃO', f"Alterado usuário '{nome}' (CPF: {cpf}).")
        else:
            senha_hash = generate_password_hash('123')
            cur = conn.execute('INSERT INTO usuarios (nome, cpf, cargo, senha_hash, primeiro_acesso, ativo, permissoes) VALUES (?, ?, ?, ?, 1, ?, ?)', (nome, cpf, cargo, senha_hash, ativo, perms_json))
            novo_id = cur.lastrowid
            for un_id in unidades_ids:
                conn.execute('INSERT INTO usuario_unidades (usuario_id, unidade_id) VALUES (?, ?)', (novo_id, int(un_id)))
            conn.commit()
            registrar_log('usuarios', 'INCLUSÃO', f"Cadastrado usuário '{nome}' (CPF: {cpf}).")
    except sqlite3.IntegrityError:
        conn.close()
        return jsonify({'status': 'erro', 'mensagem': 'Já existe um usuário cadastrado com este CPF.'})
    conn.close()
    return jsonify({'status': 'sucesso', 'mensagem': 'Usuário salvo com sucesso!'})

@app.route('/api/usuario/resetar-senha', methods=['POST'])
@login_obrigatorio
@requer_permissao('cad_usuario')
def resetar_senha_usuario():
    uid_alvo = request.form.get('id', type=int)
    uid_logada = session['unidade_id']
    eh_multi = usuario_eh_multiempresa()

    conn = get_db()
    u = conn.execute('SELECT nome, cpf, permissoes FROM usuarios WHERE id = ?', (uid_alvo,)).fetchone()
    if not u:
        conn.close()
        return jsonify({'status': 'erro', 'mensagem': 'Usuário não encontrado.'})

    if not eh_multi:
        alvo_perms = json.loads(u['permissoes'] or '[]')
        pertence = conn.execute(
            'SELECT 1 FROM usuario_unidades WHERE usuario_id = ? AND unidade_id = ?',
            (uid_alvo, uid_logada)
        ).fetchone()
        if not pertence or 'multiempresa' in alvo_perms:
            conn.close()
            return jsonify({'status': 'erro', 'mensagem': 'Você não tem permissão para resetar a senha deste usuário.'}), 403

    senha_padrao = generate_password_hash('123')
    conn.execute('UPDATE usuarios SET senha_hash = ?, primeiro_acesso = 1 WHERE id = ?', (senha_padrao, uid_alvo))
    conn.commit()
    conn.close()
    registrar_log('usuarios', 'ALTERAÇÃO', f"Resetada a senha do usuário '{u['nome']}'.")
    return jsonify({'status': 'sucesso', 'mensagem': f"Senha de {u['nome']} redefinida para '123'."})

@app.route('/api/unidade/salvar', methods=['POST'])
@login_obrigatorio
@requer_permissao('cad_unidade')
def salvar_unidade():
    id_unid = request.form.get('id', type=int)
    descricao = request.form.get('descricao', '').strip()
    ativo = request.form.get('ativo', type=int, default=1)
    eh_multi = usuario_eh_multiempresa()
    uid_logada = session['unidade_id']

    if not descricao:
        return jsonify({'status': 'erro', 'mensagem': 'Informe a descrição da unidade.'})

    if not eh_multi:
        if not id_unid or id_unid != uid_logada:
            return jsonify({'status': 'erro', 'mensagem': 'Apenas usuários com a permissão MultiEmpresa podem criar ou editar outras unidades.'}), 403

    conn = get_db()
    if id_unid:
        conn.execute('UPDATE unidades SET descricao=?, ativo=? WHERE id=?', (descricao, ativo, id_unid))
        if id_unid == uid_logada:
            session['unidade_nome'] = descricao
        conn.commit()
        registrar_log('unidades', 'ALTERAÇÃO', f"Alterada unidade '{descricao}'.")
    else:
        cur = conn.execute('INSERT INTO unidades (descricao, ativo) VALUES (?, ?)', (descricao, ativo))
        novo_id = cur.lastrowid
        conn.execute("INSERT INTO configuracoes (unidade_id, tema, logo_path) VALUES (?, 'tema_1', '')", (novo_id,))
        conn.execute('INSERT OR IGNORE INTO usuario_unidades (usuario_id, unidade_id) VALUES (?, ?)', (session['usuario_id'], novo_id))
        conn.commit()
        registrar_log('unidades', 'INCLUSÃO', f"Criada unidade '{descricao}'.")
    conn.close()
    return jsonify({'status': 'sucesso', 'mensagem': 'Unidade salva com sucesso!'})

@app.route('/api/grupo/salvar', methods=['POST'])
@login_obrigatorio
@requer_permissao('cad_grupo')
def salvar_grupo():
    id_grp = request.form.get('id', type=int)
    descricao = request.form.get('descricao', '').strip()
    ativo = request.form.get('ativo', type=int, default=1)
    refeicoes_ids = request.form.getlist('refeicoes_ids')
    uid = session['unidade_id']

    if not descricao:
        return jsonify({'status': 'erro', 'mensagem': 'Informe a descrição do grupo.'})

    conn = get_db()
    if id_grp:
        pertence = conn.execute('SELECT 1 FROM grupos WHERE id=? AND unidade_id=?', (id_grp, uid)).fetchone()
        if not pertence:
            conn.close()
            return jsonify({'status': 'erro', 'mensagem': 'Grupo não pertence a esta unidade.'}), 403
        conn.execute('UPDATE grupos SET descricao=?, ativo=? WHERE id=? AND unidade_id=?', (descricao, ativo, id_grp, uid))
        conn.execute('DELETE FROM refeicao_grupos WHERE grupo_id=?', (id_grp,))
        for r_id in refeicoes_ids:
            ref_valida = conn.execute('SELECT 1 FROM refeicoes WHERE id=? AND unidade_id=?', (int(r_id), uid)).fetchone()
            if ref_valida:
                conn.execute('INSERT INTO refeicao_grupos (refeicao_id, grupo_id) VALUES (?, ?)', (int(r_id), id_grp))
        conn.commit()
        registrar_log('grupos', 'ALTERAÇÃO', f"Alterado grupo '{descricao}'.")
    else:
        cur = conn.execute('INSERT INTO grupos (descricao, ativo, unidade_id) VALUES (?, ?, ?)', (descricao, ativo, uid))
        novo_grp = cur.lastrowid
        for r_id in refeicoes_ids:
            ref_valida = conn.execute('SELECT 1 FROM refeicoes WHERE id=? AND unidade_id=?', (int(r_id), uid)).fetchone()
            if ref_valida:
                conn.execute('INSERT INTO refeicao_grupos (refeicao_id, grupo_id) VALUES (?, ?)', (int(r_id), novo_grp))
        conn.commit()
        registrar_log('grupos', 'INCLUSÃO', f"Cadastrado grupo '{descricao}'.")
    conn.close()
    return jsonify({'status': 'sucesso', 'mensagem': 'Grupo salvo com sucesso!'})

@app.route('/api/refeicao/salvar', methods=['POST'])
@login_obrigatorio
@requer_permissao('cad_refeicao')
def salvar_refeicao():
    id_ref = request.form.get('id', type=int)
    descricao = request.form.get('descricao', '').strip()
    hora_inicio = request.form.get('hora_inicio', '').strip()
    hora_fim = request.form.get('hora_fim', '').strip()
    custo = request.form.get('custo', '0').replace(',', '.')
    ativo = request.form.get('ativo', type=int, default=1)
    grupos_ids = request.form.getlist('grupos_ids')
    uid = session['unidade_id']

    if not descricao or not hora_inicio or not hora_fim:
        return jsonify({'status': 'erro', 'mensagem': 'Preencha Descrição, Horário Inicial e Final.'})

    try:
        custo_float = float(custo)
    except ValueError:
        return jsonify({'status': 'erro', 'mensagem': 'Valor de custo inválido.'})

    conn = get_db()
    if id_ref:
        pertence = conn.execute('SELECT 1 FROM refeicoes WHERE id=? AND unidade_id=?', (id_ref, uid)).fetchone()
        if not pertence:
            conn.close()
            return jsonify({'status': 'erro', 'mensagem': 'Refeição não pertence a esta unidade.'}), 403
        conn.execute('UPDATE refeicoes SET descricao=?, hora_inicio=?, hora_fim=?, custo=?, ativo=? WHERE id=? AND unidade_id=?', (descricao, hora_inicio, hora_fim, custo_float, ativo, id_ref, uid))
        conn.execute('DELETE FROM refeicao_grupos WHERE refeicao_id=?', (id_ref,))
        for g_id in grupos_ids:
            grp_valido = conn.execute('SELECT 1 FROM grupos WHERE id=? AND unidade_id=?', (int(g_id), uid)).fetchone()
            if grp_valido:
                conn.execute('INSERT INTO refeicao_grupos (refeicao_id, grupo_id) VALUES (?, ?)', (id_ref, int(g_id)))
        conn.commit()
        registrar_log('refeicoes', 'ALTERAÇÃO', f"Alterada refeição '{descricao}'.")
    else:
        cur = conn.execute('INSERT INTO refeicoes (descricao, hora_inicio, hora_fim, custo, ativo, unidade_id) VALUES (?, ?, ?, ?, ?, ?)', (descricao, hora_inicio, hora_fim, custo_float, ativo, uid))
        novo_ref = cur.lastrowid
        for g_id in grupos_ids:
            grp_valido = conn.execute('SELECT 1 FROM grupos WHERE id=? AND unidade_id=?', (int(g_id), uid)).fetchone()
            if grp_valido:
                conn.execute('INSERT INTO refeicao_grupos (refeicao_id, grupo_id) VALUES (?, ?)', (novo_ref, int(g_id)))
        conn.commit()
        registrar_log('refeicoes', 'INCLUSÃO', f"Cadastrada refeição '{descricao}'.")
    conn.close()
    return jsonify({'status': 'sucesso', 'mensagem': 'Refeição salva com sucesso!'})

@app.route('/api/colaborador/salvar', methods=['POST'])
@login_obrigatorio
@requer_permissao('cad_colaborador')
def salvar_colaborador():
    id_col = request.form.get('id', type=int)
    nome = request.form.get('nome', '').strip()
    cpf = limpar_cpf_cnpj(request.form.get('cpf'))
    grupo_id = request.form.get('grupo_id', type=int)
    ativo = request.form.get('ativo', type=int, default=1)
    fotos_burst_json = request.form.get('fotos_burst_json', '')
    foto_b64_unica = request.form.get('foto_b64', '')
    uid = session['unidade_id']

    if not nome or len(cpf) != 11 or not grupo_id:
        return jsonify({'status': 'erro', 'mensagem': 'Preencha Nome, CPF válido (11 dígitos) e selecione um Grupo.'})

    lista_b64 = []
    if fotos_burst_json:
        try:
            lista_b64 = json.loads(fotos_burst_json)
        except Exception:
            lista_b64 = []
    elif foto_b64_unica and foto_b64_unica.startswith('data:image'):
        lista_b64 = [foto_b64_unica]

    foto_rel_path = None
    novos_vetores = []

    if lista_b64:
        for b64_item in lista_b64:
            img_bgr = decodificar_base64_para_bgr(b64_item)
            if img_bgr is None:
                continue

            if foto_rel_path is None:
                nome_seguro = f"colab_{uid}_{cpf}_{uuid.uuid4().hex[:8]}.jpg"
                caminho_completo = os.path.join(UPLOAD_FOLDER, nome_seguro)
                cv2.imwrite(caminho_completo, img_bgr)
                foto_rel_path = f"uploads/{nome_seguro}"

            vetor, _, _ = analisar_rosto_bgr(img_bgr)
            if vetor:
                novos_vetores.append(vetor)

        if not novos_vetores:
            return jsonify({
                'status': 'erro',
                'mensagem': 'Não foi possível detectar um rosto nítido nas fotos capturadas. Aproxime o rosto da câmera e tente novamente.'
            })

    conn = get_db()
    grp_valido = conn.execute('SELECT 1 FROM grupos WHERE id=? AND unidade_id=?', (grupo_id, uid)).fetchone()
    if not grp_valido:
        conn.close()
        return jsonify({'status': 'erro', 'mensagem': 'Grupo inválido para esta unidade.'}), 403

    try:
        if id_col:
            if foto_rel_path and novos_vetores:
                conn.execute('''
                    UPDATE colaboradores SET nome=?, cpf=?, foto_path=?, face_embeddings=?, grupo_id=?, ativo=?
                    WHERE id=? AND unidade_id=?
                ''', (nome, cpf, foto_rel_path, json.dumps(novos_vetores), grupo_id, ativo, id_col, uid))
            else:
                conn.execute('''
                    UPDATE colaboradores SET nome=?, cpf=?, grupo_id=?, ativo=?
                    WHERE id=? AND unidade_id=?
                ''', (nome, cpf, grupo_id, ativo, id_col, uid))
            conn.commit()
            registrar_log('colaboradores', 'ALTERAÇÃO', f"Alterado colaborador '{nome}' (CPF: {cpf}).")
        else:
            conn.execute('''
                INSERT INTO colaboradores (nome, cpf, foto_path, face_embeddings, grupo_id, ativo, unidade_id)
                VALUES (?, ?, ?, ?, ?, ?, ?)
            ''', (nome, cpf, foto_rel_path or '', json.dumps(novos_vetores), grupo_id, ativo, uid))
            conn.commit()
            registrar_log('colaboradores', 'INCLUSÃO', f"Cadastrado colaborador '{nome}' (CPF: {cpf}).")
    except sqlite3.IntegrityError:
        conn.close()
        return jsonify({'status': 'erro', 'mensagem': 'Já existe um colaborador com este CPF cadastrado nesta unidade.'})
    conn.close()

    msg_extra = f" ({len(novos_vetores)} amostras faciais gravadas com sucesso!)" if novos_vetores else ""
    return jsonify({'status': 'sucesso', 'mensagem': f"Colaborador salvo com sucesso!{msg_extra}"})

@app.route('/api/cardapio/salvar', methods=['POST'])
@login_obrigatorio
@requer_permissao('cad_cardapio')
def salvar_cardapio():
    id_card = request.form.get('id', type=int)
    data_vigencia = request.form.get('data_vigencia', '').strip()
    refeicoes_ids_raw = request.form.getlist('refeicoes_ids')
    guarnicao = request.form.get('guarnicao', '').strip()
    proteinas = request.form.get('proteinas', '').strip()
    saladas = request.form.get('saladas', '').strip()
    sobremesas = request.form.get('sobremesas', '').strip()
    bebidas = request.form.get('bebidas', '').strip()
    ativo = request.form.get('ativo', type=int, default=1)
    forcar_inativacao = request.form.get('forcar_inativacao', '0') == '1'
    uid = session['unidade_id']

    if not all([data_vigencia, guarnicao, proteinas, saladas, sobremesas, bebidas]) or not refeicoes_ids_raw:
        return jsonify({'status': 'erro', 'mensagem': 'Preencha todos os itens do cardápio, a data de vigência e selecione ao menos uma Refeição.'})

    try:
        refeicoes_ids = [int(x) for x in refeicoes_ids_raw]
    except ValueError:
        return jsonify({'status': 'erro', 'mensagem': 'Seleção de refeições inválida.'})

    conn = get_db()
    placeholders = ','.join(['?'] * len(refeicoes_ids))
    refs_validas = conn.execute(
        f'SELECT id, descricao, hora_inicio, hora_fim FROM refeicoes WHERE unidade_id = ? AND id IN ({placeholders}) ORDER BY hora_inicio',
        [uid] + refeicoes_ids
    ).fetchall()

    if len(refs_validas) != len(refeicoes_ids):
        conn.close()
        return jsonify({'status': 'erro', 'mensagem': 'Uma ou mais refeições selecionadas não pertencem a esta unidade.'}), 403

    if ativo == 1:
        query_conflito = f'''
            SELECT DISTINCT c.id, r.descricao as refeicao_nome, c.proteinas
            FROM cardapios c
            JOIN cardapio_refeicoes cr ON cr.cardapio_id = c.id
            JOIN refeicoes r ON r.id = cr.refeicao_id
            WHERE c.unidade_id = ? AND c.data_vigencia = ? AND c.ativo = 1
            AND cr.refeicao_id IN ({placeholders})
        '''
        params_conflito = [uid, data_vigencia] + refeicoes_ids
        if id_card:
            query_conflito += ' AND c.id != ?'
            params_conflito.append(id_card)

        conflitos = conn.execute(query_conflito, params_conflito).fetchall()
        if conflitos and not forcar_inativacao:
            conn.close()
            desc_conflito = ", ".join([f"Cardápio #{c['id']} ({c['refeicao_nome']} - {c['proteinas']})" for c in conflitos])
            return jsonify({
                'status': 'conflito',
                'mensagem': f"Já existe outro cardápio ativo na data {data_vigencia} vinculado à mesma refeição: {desc_conflito}. Deseja inativar o outro cardápio para ativar este?"
            })
        elif conflitos and forcar_inativacao:
            ids_inativados = set()
            for c in conflitos:
                if c['id'] not in ids_inativados:
                    conn.execute('UPDATE cardapios SET ativo = 0 WHERE id = ? AND unidade_id = ?', (c['id'], uid))
                    ids_inativados.add(c['id'])
                    registrar_log('cardapios', 'ALTERAÇÃO', f"Cardápio #{c['id']} inativado automaticamente por conflito de vigência na data {data_vigencia}.")

    hora_ini_resumo = refs_validas[0]['hora_inicio']
    hora_fim_resumo = refs_validas[-1]['hora_fim']
    nomes_refs = ", ".join([r['descricao'] for r in refs_validas])

    if id_card:
        pertence = conn.execute('SELECT 1 FROM cardapios WHERE id = ? AND unidade_id = ?', (id_card, uid)).fetchone()
        if not pertence:
            conn.close()
            return jsonify({'status': 'erro', 'mensagem': 'Cardápio não pertence a esta unidade.'}), 403

        conn.execute('''
            UPDATE cardapios
            SET data_vigencia=?, hora_inicio=?, hora_fim=?, guarnicao=?, proteinas=?, saladas=?, sobremesas=?, bebidas=?, ativo=?
            WHERE id=? AND unidade_id=?
        ''', (data_vigencia, hora_ini_resumo, hora_fim_resumo, guarnicao, proteinas, saladas, sobremesas, bebidas, ativo, id_card, uid))
        conn.execute('DELETE FROM cardapio_refeicoes WHERE cardapio_id = ?', (id_card,))
        for r_id in refeicoes_ids:
            conn.execute('INSERT INTO cardapio_refeicoes (cardapio_id, refeicao_id) VALUES (?, ?)', (id_card, r_id))
        conn.commit()
        registrar_log('cardapios', 'ALTERAÇÃO', f"Alterado cardápio de {data_vigencia} vinculado a: {nomes_refs}.")
    else:
        cur = conn.execute('''
            INSERT INTO cardapios (data_vigencia, hora_inicio, hora_fim, guarnicao, proteinas, saladas, sobremesas, bebidas, ativo, unidade_id)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (data_vigencia, hora_ini_resumo, hora_fim_resumo, guarnicao, proteinas, saladas, sobremesas, bebidas, ativo, uid))
        novo_card_id = cur.lastrowid
        for r_id in refeicoes_ids:
            conn.execute('INSERT INTO cardapio_refeicoes (cardapio_id, refeicao_id) VALUES (?, ?)', (novo_card_id, r_id))
        conn.commit()
        registrar_log('cardapios', 'INCLUSÃO', f"Cadastrado cardápio para {data_vigencia} vinculado a: {nomes_refs}.")

    conn.close()
    return jsonify({'status': 'sucesso', 'mensagem': 'Cardápio salvo com sucesso!'})

@app.route('/api/configuracoes/salvar', methods=['POST'])
@login_obrigatorio
@requer_permissao('configuracoes')
def salvar_configuracoes():
    tema = request.form.get('tema', 'tema_1')
    uid = session['unidade_id']
    if tema not in TEMAS:
        tema = 'tema_1'

    logo_file = request.files.get('logo')
    logo_rel_path = None
    if logo_file and logo_file.filename:
        if not arquivo_permitido(logo_file.filename):
            return jsonify({'status': 'erro', 'mensagem': 'Formato de logo inválido.'})
        ext = logo_file.filename.rsplit('.', 1)[1].lower()
        nome_seguro = secure_filename(f"logo_unidade_{uid}_{uuid.uuid4().hex[:6]}.{ext}")
        logo_file.save(os.path.join(UPLOAD_FOLDER, nome_seguro))
        logo_rel_path = f"uploads/{nome_seguro}"

    conn = get_db()
    existe = conn.execute('SELECT unidade_id FROM configuracoes WHERE unidade_id = ?', (uid,)).fetchone()
    if existe:
        if logo_rel_path:
            conn.execute('UPDATE configuracoes SET tema=?, logo_path=? WHERE unidade_id=?', (tema, logo_rel_path, uid))
        else:
            conn.execute('UPDATE configuracoes SET tema=? WHERE unidade_id=?', (tema, uid))
    else:
        conn.execute('INSERT INTO configuracoes (unidade_id, tema, logo_path) VALUES (?, ?, ?)', (uid, tema, logo_rel_path or ''))
    conn.commit()
    conn.close()
    registrar_log('configuracoes', 'ALTERAÇÃO', "Atualizadas configurações da unidade.")
    return jsonify({'status': 'sucesso', 'mensagem': 'Configurações de Tema e Logo aplicadas!'})

@app.route('/api/logs/consultar', methods=['POST'])
@login_obrigatorio
@requer_permissao('logs')
def consultar_logs():
    data_ini = request.form.get('data_inicio', '')
    data_fim = request.form.get('data_fim', '')
    tabela = request.form.get('tabela', 'TODAS')
    uid = session['unidade_id']

    query = '''
        SELECT data_hora, usuario_nome, tabela, acao, descricao FROM logs
        WHERE unidade_id = ?
        AND substr(data_hora, 1, 10) >= ? AND substr(data_hora, 1, 10) <= ?
    '''
    params = [uid, data_ini, data_fim]
    if tabela != 'TODAS':
        query += ' AND tabela = ?'
        params.append(tabela)

    query += ' ORDER BY id DESC LIMIT 500'
    conn = get_db()
    rows = conn.execute(query, params).fetchall()
    conn.close()
    return jsonify({'status': 'sucesso', 'logs': [dict(r) for r in rows]})

# ==============================================================================
# RELATÓRIOS EM PDF
# ==============================================================================
class RelatorioPDF(FPDF):
    def __init__(self, titulo, unidade_nome, periodo_str, logo_path):
        super().__init__()
        self.titulo = titulo
        self.unidade_nome = unidade_nome
        self.periodo_str = periodo_str
        self.logo_path = logo_path

    def header(self):
        if self.logo_path:
            caminho_abs = os.path.join(BASE_DIR, 'static', self.logo_path)
            if os.path.exists(caminho_abs):
                try:
                    self.image(caminho_abs, 10, 8, 25)
                except Exception:
                    pass
        self.set_font('Helvetica', 'B', 14)
        self.cell(0, 8, self.titulo.encode('latin-1', 'replace').decode('latin-1'), ln=True, align='C')
        self.set_font('Helvetica', '', 10)
        sub = f"Unidade: {self.unidade_nome}   |   Período: {self.periodo_str}"
        self.cell(0, 6, sub.encode('latin-1', 'replace').decode('latin-1'), ln=True, align='C')
        self.ln(6)
        self.line(10, self.get_y(), 200, self.get_y())
        self.ln(4)

    def footer(self):
        self.set_y(-15)
        self.set_font('Helvetica', 'I', 8)
        rodape = f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} - Página {self.page_no()}/{{nb}}"
        self.cell(0, 10, rodape.encode('latin-1', 'replace').decode('latin-1'), align='C')

def txt_pdf(texto):
    return str(texto or '').encode('latin-1', 'replace').decode('latin-1')

@app.route('/admin/relatorios/pdf', methods=['GET'])
@login_obrigatorio
@requer_permissao('relatorios')
def gerar_relatorio_pdf():
    tipo = request.args.get('tipo', 'analitico')
    data_ini = request.args.get('data_inicio', datetime.now().strftime('%Y-%m-01'))
    data_fim = request.args.get('data_fim', datetime.now().strftime('%Y-%m-%d'))
    uid = session['unidade_id']
    cfg = obter_config_unidade(uid)
    periodo_fmt = f"{data_ini} a {data_fim}"
    conn = get_db()

    if tipo == 'analitico':
        pdf = RelatorioPDF("Relatório Analítico de Refeições por Colaborador", cfg['unidade_nome'], periodo_fmt, cfg['logo_path'])
        pdf.alias_nb_pages()
        pdf.add_page()
        rows = conn.execute('''
            SELECT rr.data_registro, rr.hora_registro, c.nome as colab_nome, c.cpf,
                   r.descricao as refeicao_nome, rr.custo_registrado, rr.metodo
            FROM registros_refeicao rr
            JOIN colaboradores c ON c.id = rr.colaborador_id
            JOIN refeicoes r ON r.id = rr.refeicao_id
            WHERE rr.unidade_id = ? AND rr.data_registro >= ? AND rr.data_registro <= ?
            ORDER BY c.nome ASC, rr.data_registro ASC, rr.hora_registro ASC
        ''', (uid, data_ini, data_fim)).fetchall()

        pdf.set_font('Helvetica', 'B', 9)
        pdf.set_fill_color(230, 230, 230)
        pdf.cell(25, 8, txt_pdf("Data/Hora"), 1, 0, 'C', True)
        pdf.cell(65, 8, txt_pdf("Colaborador"), 1, 0, 'L', True)
        pdf.cell(30, 8, txt_pdf("CPF"), 1, 0, 'C', True)
        pdf.cell(45, 8, txt_pdf("Refeição"), 1, 0, 'L', True)
        pdf.cell(25, 8, txt_pdf("Custo (R$)"), 1, 1, 'R', True)

        pdf.set_font('Helvetica', '', 9)
        total_geral = 0.0
        totais_colaborador = {}
        for r in rows:
            dt_hr = f"{r['data_registro']} {r['hora_registro'][:5]}"
            pdf.cell(25, 7, txt_pdf(dt_hr), 1, 0, 'C')
            pdf.cell(65, 7, txt_pdf(r['colab_nome'][:34]), 1, 0, 'L')
            pdf.cell(30, 7, txt_pdf(r['cpf']), 1, 0, 'C')
            pdf.cell(45, 7, txt_pdf(r['refeicao_nome'][:24]), 1, 0, 'L')
            pdf.cell(25, 7, txt_pdf(f"R$ {r['custo_registrado']:.2f}"), 1, 1, 'R')
            total_geral += r['custo_registrado']
            chave = f"{r['colab_nome']} ({r['cpf']})"
            totais_colaborador.setdefault(chave, {'qtd': 0, 'soma': 0.0})
            totais_colaborador[chave]['qtd'] += 1
            totais_colaborador[chave]['soma'] += r['custo_registrado']

        pdf.ln(6)
        pdf.set_font('Helvetica', 'B', 10)
        pdf.cell(0, 8, txt_pdf("Resumo de Custo Acumulado por Colaborador:"), ln=True)
        pdf.set_font('Helvetica', 'B', 9)
        pdf.cell(120, 7, txt_pdf("Colaborador"), 1, 0, 'L', True)
        pdf.cell(30, 7, txt_pdf("Qtd. Refeições"), 1, 0, 'C', True)
        pdf.cell(40, 7, txt_pdf("Custo Gerado"), 1, 1, 'R', True)

        pdf.set_font('Helvetica', '', 9)
        for colab, dados in totais_colaborador.items():
            pdf.cell(120, 7, txt_pdf(colab[:60]), 1, 0, 'L')
            pdf.cell(30, 7, str(dados['qtd']), 1, 0, 'C')
            pdf.cell(40, 7, txt_pdf(f"R$ {dados['soma']:.2f}"), 1, 1, 'R')

        pdf.set_font('Helvetica', 'B', 10)
        pdf.cell(150, 8, txt_pdf(f"TOTAL GERAL ({len(rows)} refeições):"), 1, 0, 'R', True)
        pdf.cell(40, 8, txt_pdf(f"R$ {total_geral:.2f}"), 1, 1, 'R', True)

    elif tipo == 'sintetico':
        pdf = RelatorioPDF("Relatório Sintético de Refeições", cfg['unidade_nome'], periodo_fmt, cfg['logo_path'])
        pdf.alias_nb_pages()
        pdf.add_page()
        rows = conn.execute('''
            SELECT r.descricao as refeicao_nome, COUNT(rr.id) as quantidade, COALESCE(SUM(rr.custo_registrado), 0) as custo_total
            FROM registros_refeicao rr
            JOIN refeicoes r ON r.id = rr.refeicao_id
            WHERE rr.unidade_id = ? AND rr.data_registro >= ? AND rr.data_registro <= ?
            GROUP BY r.id, r.descricao
            ORDER BY r.descricao ASC
        ''', (uid, data_ini, data_fim)).fetchall()

        pdf.set_font('Helvetica', 'B', 10)
        pdf.set_fill_color(230, 230, 230)
        pdf.cell(95, 8, txt_pdf("Tipo de Refeição"), 1, 0, 'L', True)
        pdf.cell(45, 8, txt_pdf("Total de Refeições"), 1, 0, 'C', True)
        pdf.cell(50, 8, txt_pdf("Custo Total (R$)"), 1, 1, 'R', True)

        pdf.set_font('Helvetica', '', 10)
        soma_qtd = 0
        soma_custo = 0.0
        for r in rows:
            pdf.cell(95, 8, txt_pdf(r['refeicao_nome']), 1, 0, 'L')
            pdf.cell(45, 8, str(r['quantidade']), 1, 0, 'C')
            pdf.cell(50, 8, txt_pdf(f"R$ {r['custo_total']:.2f}"), 1, 1, 'R')
            soma_qtd += r['quantidade']
            soma_custo += r['custo_total']

        pdf.set_font('Helvetica', 'B', 10)
        pdf.cell(95, 9, txt_pdf("TOTAL CONSOLIDADO DA UNIDADE"), 1, 0, 'R', True)
        pdf.cell(45, 9, str(soma_qtd), 1, 0, 'C', True)
        pdf.cell(50, 9, txt_pdf(f"R$ {soma_custo:.2f}"), 1, 1, 'R', True)

    else:
        pdf = RelatorioPDF("Relatório de Cardápios Ofertados", cfg['unidade_nome'], periodo_fmt, cfg['logo_path'])
        pdf.alias_nb_pages()
        pdf.add_page()
        rows = conn.execute('''
            SELECT * FROM cardapios WHERE unidade_id = ? AND data_vigencia >= ? AND data_vigencia <= ?
            ORDER BY data_vigencia ASC, id ASC
        ''', (uid, data_ini, data_fim)).fetchall()

        for c in rows:
            refs_c = conn.execute('''
                SELECT r.descricao, r.hora_inicio, r.hora_fim
                FROM cardapio_refeicoes cr
                JOIN refeicoes r ON r.id = cr.refeicao_id
                WHERE cr.cardapio_id = ?
                ORDER BY r.hora_inicio ASC
            ''', (c['id'],)).fetchall()

            if refs_c:
                refs_str = ", ".join([f"{r['descricao']} ({r['hora_inicio']}-{r['hora_fim']})" for r in refs_c])
            else:
                refs_str = f"{c['hora_inicio']} às {c['hora_fim']}" if c['hora_inicio'] else "Sem refeição vinculada"

            pdf.set_font('Helvetica', 'B', 10)
            pdf.set_fill_color(230, 230, 230)
            status_str = "Ativo" if c['ativo'] == 1 else "Inativo"
            pdf.cell(0, 8, txt_pdf(f"Data: {c['data_vigencia']} | Refeições: {refs_str} | Status: {status_str}"), 1, 1, 'L', True)
            pdf.set_font('Helvetica', '', 9)
            pdf.multi_cell(0, 6, txt_pdf(f"Proteínas: {c['proteinas']}\nGuarnição: {c['guarnicao']}\nSaladas: {c['saladas']}\nSobremesas: {c['sobremesas']}\nBebidas: {c['bebidas']}"), 1)
            pdf.ln(3)

    conn.close()
    caminho_pdf = os.path.join(UPLOAD_FOLDER, f"relatorio_{tipo}_{uid}.pdf")
    pdf.output(caminho_pdf)
    registrar_log('relatorios', 'CONSULTA', f"Emitido relatório PDF ({tipo}) de {data_ini} até {data_fim}.")
    return send_file(caminho_pdf, as_attachment=True, download_name=f"relatorio_{tipo}_{data_ini}_{data_fim}.pdf")

# ==============================================================================
# MÓDULO 2: TOTEM DE ENTRADA
# ==============================================================================
@app.route('/totem')
def totem():
    conn = get_db()
    totem_uid = session.get('totem_unidade_id')
    qtd_rostos_ativos = 0
    if totem_uid:
        cols = conn.execute(
            "SELECT face_embeddings FROM colaboradores WHERE unidade_id = ? AND ativo = 1 AND face_embeddings != '[]' AND face_embeddings IS NOT NULL",
            (totem_uid,)
        ).fetchall()
        qtd_rostos_ativos = len(cols)
    conn.close()
    return render_template('totem.html', totem_uid=totem_uid, qtd_rostos_ativos=qtd_rostos_ativos)

@app.route('/api/totem/autenticar-unidade', methods=['POST'])
def totem_autenticar_unidade():
    unidade_id = request.form.get('unidade_id', type=int)
    cpf = limpar_cpf_cnpj(request.form.get('cpf'))
    senha = request.form.get('senha', '')

    if not cpf or not senha or not unidade_id:
        return jsonify({'status': 'erro', 'mensagem': 'Informe CPF, senha e selecione a unidade.'})

    conn = get_db()
    user = conn.execute('SELECT * FROM usuarios WHERE cpf = ? AND ativo = 1', (cpf,)).fetchone()
    if not user or not check_password_hash(user['senha_hash'], senha):
        conn.close()
        return jsonify({'status': 'erro', 'mensagem': 'Credenciais inválidas para liberar o Totem.'})

    perms = json.loads(user['permissoes'] or '[]')
    eh_multi = 'multiempresa' in perms

    unid = conn.execute('SELECT descricao FROM unidades WHERE id = ? AND ativo = 1', (unidade_id,)).fetchone()
    if not unid:
        conn.close()
        return jsonify({'status': 'erro', 'mensagem': 'Unidade inválida ou inativa.'})

    if not eh_multi:
        vinculo = conn.execute('SELECT 1 FROM usuario_unidades WHERE usuario_id = ? AND unidade_id = ?', (user['id'], unidade_id)).fetchone()
        if not vinculo:
            conn.close()
            return jsonify({'status': 'erro', 'mensagem': 'Este usuário não possui acesso à unidade selecionada.'})

    conn.close()

    session['totem_unidade_id'] = unidade_id
    session['totem_unidade_nome'] = unid['descricao']
    registrar_log('unidades', 'TOTEM', f"Totem vinculado à unidade '{unid['descricao']}' por '{user['nome']}'.", unidade_id)
    return jsonify({'status': 'sucesso', 'mensagem': f"Unidade '{unid['descricao']}' autenticada no Totem!"})

@app.route('/api/totem/cardapio-atual', methods=['GET'])
def totem_cardapio_atual():
    uid = session.get('totem_unidade_id')
    if not uid:
        return jsonify({'status': 'sem_unidade'})

    agora = datetime.now()
    data_hoje = agora.strftime('%Y-%m-%d')
    hora_atual = agora.strftime('%H:%M')

    conn = get_db()
    refeicao = conn.execute('''
        SELECT id, descricao, hora_inicio, hora_fim FROM refeicoes
        WHERE unidade_id = ? AND ativo = 1
        AND hora_inicio <= ? AND hora_fim >= ?
        ORDER BY id DESC LIMIT 1
    ''', (uid, hora_atual, hora_atual)).fetchone()

    cardapio = conn.execute('''
        SELECT c.* FROM cardapios c
        JOIN cardapio_refeicoes cr ON cr.cardapio_id = c.id
        JOIN refeicoes r ON r.id = cr.refeicao_id
        WHERE c.unidade_id = ? AND c.data_vigencia = ? AND c.ativo = 1
        AND r.unidade_id = ? AND r.ativo = 1
        AND r.hora_inicio <= ? AND r.hora_fim >= ?
        ORDER BY c.id DESC LIMIT 1
    ''', (uid, data_hoje, uid, hora_atual, hora_atual)).fetchone()

    conn.close()

    return jsonify({
        'status': 'sucesso',
        'hora_servidor': hora_atual,
        'data_servidor': agora.strftime('%d/%m/%Y'),
        'cardapio': dict(cardapio) if cardapio else None,
        'refeicao': dict(refeicao) if refeicao else None
    })

def validar_refeicao_colaborador_no_horario(conn, uid, colab):
    agora = datetime.now()
    hora_atual = agora.strftime('%H:%M')
    data_hoje = agora.strftime('%Y-%m-%d')

    refeicoes_horario = conn.execute('''
        SELECT r.* FROM refeicoes r
        JOIN refeicao_grupos rg ON rg.refeicao_id = r.id
        JOIN grupos g ON g.id = rg.grupo_id
        WHERE r.unidade_id = ? AND r.ativo = 1 AND g.ativo = 1
        AND rg.grupo_id = ? AND r.hora_inicio <= ? AND r.hora_fim >= ?
        ORDER BY r.hora_inicio ASC
    ''', (uid, colab['grupo_id'], hora_atual, hora_atual)).fetchall()

    if not refeicoes_horario:
        return None, f"Olá, {colab['nome']}! Não há refeição liberada para o seu grupo neste horário ({hora_atual})."

    ultimo_registro_msg = None
    for ref in refeicoes_horario:
        ja_registrou = conn.execute('''
            SELECT hora_registro FROM registros_refeicao
            WHERE colaborador_id = ? AND refeicao_id = ? AND data_registro = ?
        ''', (colab['id'], ref['id'], data_hoje)).fetchone()

        if not ja_registrou:
            return ref, None
        else:
            ultimo_registro_msg = (
                f"{colab['nome']}, sua entrada para '{ref['descricao']}' já foi registrada hoje às "
                f"{ja_registrou['hora_registro'][:5]}. É permitida apenas 1 entrada por refeição no dia."
            )

    return None, ultimo_registro_msg

@app.route('/api/totem/reconhecer-frame', methods=['POST'])
def totem_reconhecer_frame():
    uid = session.get('totem_unidade_id')
    if not uid:
        return jsonify({'status': 'sem_unidade'})

    frame_b64 = request.form.get('frame_b64', '')
    img_bgr = decodificar_base64_para_bgr(frame_b64)
    if img_bgr is None:
        return jsonify({'status': 'sem_rosto'})

    vetor, bbox, proporcao = analisar_rosto_bgr(img_bgr)
    if vetor is None or bbox is None:
        return jsonify({'status': 'sem_rosto'})

    if proporcao < 0.045:
        return jsonify({'status': 'longe', 'bbox': bbox})

    conn = get_db()
    colabs = conn.execute('''
        SELECT id, nome, cpf, foto_path, grupo_id, face_embeddings
        FROM colaboradores
        WHERE unidade_id = ? AND ativo = 1 AND face_embeddings IS NOT NULL AND face_embeddings != '[]'
    ''', (uid,)).fetchall()

    melhor_colab = None
    melhor_score = 0.0

    for c in colabs:
        try:
            lista_emb = json.loads(c['face_embeddings'] or '[]')
            score = comparar_vetor_com_lista(vetor, lista_emb)
            if score > melhor_score:
                melhor_score = score
                melhor_colab = c
        except Exception:
            continue

    if not melhor_colab or melhor_score < LIMIAR_SIMILARIDADE:
        conn.close()
        return jsonify({
            'status': 'desconhecido',
            'bbox': bbox,
            'score': round(melhor_score, 2)
        })

    refeicao, erro_msg = validar_refeicao_colaborador_no_horario(conn, uid, melhor_colab)
    conn.close()

    if not refeicao:
        return jsonify({
            'status': 'bloqueado',
            'bbox': bbox,
            'nome': melhor_colab['nome'],
            'mensagem': erro_msg
        })

    return jsonify({
        'status': 'reconhecido',
        'bbox': bbox,
        'score': round(melhor_score, 2),
        'colaborador': {
            'id': melhor_colab['id'],
            'nome': melhor_colab['nome'],
            'cpf': melhor_colab['cpf'],
            'foto_path': url_for('static', filename=melhor_colab['foto_path']) if melhor_colab['foto_path'] else ''
        },
        'refeicao': {
            'id': refeicao['id'],
            'descricao': refeicao['descricao'],
            'custo': refeicao['custo']
        },
        'metodo': 'FACIAL'
    })

@app.route('/api/totem/identificar', methods=['POST'])
def totem_identificar():
    uid = session.get('totem_unidade_id')
    if not uid:
        return jsonify({'status': 'erro', 'mensagem': 'Autentique uma unidade antes de registrar.'})

    cpf = limpar_cpf_cnpj(request.form.get('cpf'))
    colaborador_id = request.form.get('colaborador_id', type=int)
    metodo = request.form.get('metodo', 'CPF')

    conn = get_db()
    if colaborador_id:
        colab = conn.execute('SELECT * FROM colaboradores WHERE id = ? AND unidade_id = ?', (colaborador_id, uid)).fetchone()
    else:
        colab = conn.execute('SELECT * FROM colaboradores WHERE cpf = ? AND unidade_id = ?', (cpf, uid)).fetchone()

    if not colab:
        conn.close()
        return jsonify({'status': 'erro', 'mensagem': 'Colaborador não encontrado nesta unidade.'})

    if colab['ativo'] == 0:
        conn.close()
        return jsonify({'status': 'erro', 'mensagem': f"O cadastro de {colab['nome']} está inativo."})

    refeicao, erro_msg = validar_refeicao_colaborador_no_horario(conn, uid, colab)
    conn.close()

    if not refeicao:
        return jsonify({'status': 'erro', 'mensagem': erro_msg})

    return jsonify({
        'status': 'sucesso',
        'colaborador': {
            'id': colab['id'], 'nome': colab['nome'], 'cpf': colab['cpf'],
            'foto_path': url_for('static', filename=colab['foto_path']) if colab['foto_path'] else ''
        },
        'refeicao': {'id': refeicao['id'], 'descricao': refeicao['descricao'], 'custo': refeicao['custo']},
        'metodo': metodo
    })

@app.route('/api/totem/confirmar-entrada', methods=['POST'])
def totem_confirmar_entrada():
    uid = session.get('totem_unidade_id')
    if not uid:
        return jsonify({'status': 'erro', 'mensagem': 'Unidade não autenticada no Totem.'})

    colaborador_id = request.form.get('colaborador_id', type=int)
    refeicao_id = request.form.get('refeicao_id', type=int)
    metodo = request.form.get('metodo', 'CPF')
    frame_b64 = request.form.get('frame_b64', '')

    agora = datetime.now()
    data_hoje = agora.strftime('%Y-%m-%d')
    hora_atual = agora.strftime('%H:%M')
    hora_agora = agora.strftime('%H:%M:%S')

    conn = get_db()
    colab = conn.execute('SELECT id, nome, cpf, grupo_id FROM colaboradores WHERE id = ? AND unidade_id = ? AND ativo = 1', (colaborador_id, uid)).fetchone()
    ref = conn.execute('''
        SELECT r.id, r.descricao, r.custo
        FROM refeicoes r
        JOIN refeicao_grupos rg ON rg.refeicao_id = r.id
        JOIN grupos g ON g.id = rg.grupo_id
        WHERE r.id = ? AND r.unidade_id = ? AND r.ativo = 1 AND g.ativo = 1
        AND rg.grupo_id = ? AND r.hora_inicio <= ? AND r.hora_fim >= ?
    ''', (refeicao_id, uid, colab['grupo_id'] if colab else -1, hora_atual, hora_atual)).fetchone()

    if not colab or not ref:
        conn.close()
        return jsonify({'status': 'erro', 'mensagem': 'Horário da refeição encerrado ou dados inválidos.'})

    ja_registrou = conn.execute('''
        SELECT hora_registro FROM registros_refeicao
        WHERE colaborador_id = ? AND refeicao_id = ? AND data_registro = ?
    ''', (colaborador_id, refeicao_id, data_hoje)).fetchone()

    if ja_registrou:
        conn.close()
        return jsonify({
            'status': 'erro',
            'mensagem': f"{colab['nome']}, você já registrou entrada para '{ref['descricao']}' hoje às {ja_registrou['hora_registro'][:5]}. Só é permitida 1 entrada por refeição no dia."
        })

    try:
        conn.execute('''
            INSERT INTO registros_refeicao (colaborador_id, refeicao_id, unidade_id, data_registro, hora_registro, custo_registrado, metodo)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        ''', (colaborador_id, refeicao_id, uid, data_hoje, hora_agora, ref['custo'], metodo))
        conn.commit()
    except sqlite3.IntegrityError:
        conn.close()
        return jsonify({'status': 'erro', 'mensagem': 'Entrada já registrada para esta refeição hoje.'})

    conn.close()

    if frame_b64:
        img_bgr = decodificar_base64_para_bgr(frame_b64)
        vetor_novo, _, proporcao = analisar_rosto_bgr(img_bgr)
        if vetor_novo and proporcao >= 0.04:
            adicionar_vetor_aprendizado_continuo(colaborador_id, vetor_novo)

    registrar_log('registros_refeicao', 'INCLUSÃO', f"Entrada confirmada: '{colab['nome']}' - '{ref['descricao']}' via {metodo}.", uid)
    return jsonify({'status': 'sucesso', 'mensagem': f"Bom apetite, {colab['nome']}! Entrada em '{ref['descricao']}' registrada com sucesso."})

if __name__ == '__main__':
    app.run(host='127.0.0.1', port=5055, debug=False)